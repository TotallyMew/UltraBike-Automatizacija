from __future__ import annotations

import html
import json
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from tools.orbea_automation.catalogue import CatalogueEntry, MatchResult
from tools.orbea_automation.models import CancellationToken, OrbeaRunConfig, OrbeaRunFailure
from tools.orbea_automation.service import OrbeaAutomationService
from tools.orbea_automation.website import OrbeaAccessError, OrbeaWebsiteClient, SearchProduct, parse_website_product, product_matches


URL = "https://www.orbea.com/en-be/orca-m30i"
SKU = "U10707SV"


def page(code=SKU, name="ORCA M30i", description="A road bike for long days."):
    product = {"@type": "Product", "name": name, "sku": code, "description": description,
               "additionalProperty": [{"name": "Frame", "value": "Orbea carbon"}]}
    return f'<html><body><nav>U107TTCC</nav><main><h1>{html.escape(name)}</h1><script type="application/ld+json">{json.dumps(product)}</script></main></body></html>'


def row():
    return {"row_key": "p1", "product_id": "p1", "sku": SKU, "visible_code": "U107",
            "title": "Orbea ORCA M30i Black", "status": "unmatched", "product_url": "https://pim.bo.ultrabike.lt/dashboard/products/p1", "page": 1, "row": 1}


class Website:
    def __init__(self, product=None):
        self.product = product or parse_website_product(page(), URL)
        # Photo fixtures expose the same colour metadata as the official configurator.
        photo_template = {"name": self.product.name, "hash": "photo-fixture", "code": "U107TTCC",
            "views": [{"type": "side", "status": "published"}],
            "zones": [{"identifier": "C1", "type": "frame", "default_color": "SV", "colors": [
                {"color": {"code": "SV", "name": {"en": "Black"}, "status": "published"}}]}]}
        initializer = "currentTemplate = JSON.parse('" + json.dumps(photo_template) + "');"
        photo_html = '<div id="product-bike-detail" x-init="' + html.escape(initializer, quote=True) + '"></div>'
        self.product = replace(self.product, page_html=photo_html + self.product.page_html)
        self.lookups = []
        self.fetches = []
        self.closed = False

    def lookup(self, sku, title):
        self.lookups.append(sku)
        return MatchResult("code_match", "Orbea website code", CatalogueEntry("U107", "U107TTCC", "ORCA M30i", product_link=URL))

    def fetch(self, url):
        self.fetches.append(url)
        return self.product

    def close(self):
        self.closed = True

    cancel = close


class Photos:
    def __init__(self):
        self.calls = 0
        self.fail = False
        self.cancel_run = False

    def run_from_html(self, url, page_html, output_dir, **kwargs):
        self.calls += 1
        assert url == URL and '"sku": "U10707SV"' in page_html
        folder = Path(output_dir) / kwargs["product_folder"]
        folder.mkdir(parents=True, exist_ok=True)
        image = folder / "black-side.png"
        image.write_bytes(b"photo fixture")
        return SimpleNamespace(files=(image,), failures=("temporary asset error",) if self.fail else (), unavailable=(), cancelled=self.cancel_run)


def scan(_client, catalogue, checkpoint, config, **kwargs):
    # The existing scanner supplies one representative SKU per product.
    checkpoint.upsert_result(row())
    checkpoint.data["scan_completed"] = True
    checkpoint.save()


def tables(_service, config, checkpoint, reporter, token, log, **kwargs):
    folder = checkpoint.run_dir / "images" / "orca"
    folder.mkdir(parents=True, exist_ok=True)
    for filename in ("geometry.png", "geometry-xs.png", "size-guide-cm.png"):
        (folder / filename).write_bytes(b"table fixture")
    checkpoint.upsert_image(URL, {"folder": "images/orca", "geometry_image": "images/orca/geometry.png", "size_guide_image": "images/orca/size-guide-cm.png",
        "geometry_status": "downloaded", "size_guide_status": "downloaded", "geometry_variants": [{"filename": "geometry-xs.png"}]})
    checkpoint.data["images_completed"] = True
    checkpoint.save()


def config(root, **options):
    values = dict(collect_product_data=True, download_images=True, download_product_photos=True,
                  download_description=True, download_specifications=True, product_code_prefix="U")
    values.update(options)
    return OrbeaRunConfig(None, root / "runs", **values)


def service(website, photos):
    return OrbeaAutomationService(object(), website_client_factory=lambda *_: website, photo_service_factory=lambda: photos)


def test_product_identifiers_ignore_navigation_search_echo_and_recommendations():
    source = page("T12302AA") + '<input value="U10707SV"><script type="application/ld+json">{"@type":"Product","name":"Other Bike","sku":"U10707SV"}</script>'
    product = parse_website_product(source, URL)
    assert product.codes == ("T12302AA",)
    assert not product_matches(product, SKU, "Orbea ORCA M30i")


def test_access_check_is_not_reported_as_a_product_or_missing_assets():
    with pytest.raises(OrbeaAccessError, match="security check"):
        parse_website_product('<h1>www.orbea.com</h1><p>Verify you are human</p>', URL)
    client = OrbeaWebsiteClient(lambda: None, CancellationToken())
    client.regions = (client.HOME,)
    with patch.object(client, "_search_results", side_effect=OrbeaAccessError("access check")) as search:
        with pytest.raises(OrbeaAccessError):
            client.lookup(SKU, "ORCA M30i")
    assert search.call_count == 1


class DelayedSecurityPage:
    """Eager navigation returns before Cloudflare replaces the initial page."""

    def __init__(self, initial_reads):
        self.initial_reads = initial_reads
        self.reads = 0
        self.navigations = []
        self.current_url = URL

    def set_page_load_timeout(self, timeout):
        pass

    def get(self, url):
        self.navigations.append(url)
        self.current_url = url

    @property
    def page_source(self):
        self.reads += 1
        if self.reads <= self.initial_reads:
            return '<html><main>Loading</main></html>'
        return '<html><title>Just a moment...</title><input name="cf-turnstile-response" type="hidden"></html>'

    def find_elements(self, by, selector):
        return []


def test_security_page_that_appears_while_waiting_for_search_stops_first_query():
    driver = DelayedSecurityPage(initial_reads=2)
    client = OrbeaWebsiteClient(lambda: driver, CancellationToken(), timeout=1, access_timeout=0)
    with pytest.raises(OrbeaAccessError, match="partial Excel"):
        client.lookup(SKU, "Orbea ORCA M30i")
    assert driver.navigations == [client.HOME]
    assert driver.reads == 4


def test_security_page_that_appears_while_waiting_for_product_is_detected():
    driver = DelayedSecurityPage(initial_reads=1)
    client = OrbeaWebsiteClient(lambda: driver, CancellationToken(), timeout=1, access_timeout=0)
    with pytest.raises(OrbeaAccessError, match="security check"):
        client.fetch(URL)
    assert driver.reads == 3 and not client._products


def test_template_code_matches_selected_template_without_title_matching():
    initializer = "currentTemplate = JSON.parse('" + json.dumps({"code": "U107TTCC", "name": "ORCA M30i"}) + "');"
    source = '<main><h1>ORCA M30i</h1><div id="product-bike-detail" x-init="' + html.escape(initializer, quote=True) + '"></div></main>'
    product = parse_website_product(source, URL)
    assert product_matches(product, SKU, "Orbea ORCA M30i Black")
    assert product_matches(product, SKU, "Orbea ORCA M20i")
    assert not product_matches(product, "U10907SV", "Orbea ORCA M30i")


def test_visible_component_rows_and_description_are_extracted_without_site_scripts():
    source = '<main><h1>ORCA M30i</h1><p>A lightweight road bike for long rides and mountain climbs.</p><section id="specifications"><h3>Fork</h3><p>Orbea carbon</p><h3>Brakes</h3><p>Shimano hydraulic disc</p></section></main>'
    product = parse_website_product(source, URL)
    assert product.specifications_text == "Fork: Orbea carbon\nBrakes: Shimano hydraulic disc"
    assert "lightweight" in product.description_html and "Shimano" not in product.description_html


def test_direct_search_requires_exact_code_and_uses_first_matching_row():
    client = OrbeaWebsiteClient(lambda: None, CancellationToken())
    client.regions = (client.HOME,)
    with patch.object(client, "_search_results", return_value=(SearchProduct("T123TTCC", "ORCA M30i", URL),)):
        assert client.lookup(SKU, "Orbea ORCA M30i").status == "unmatched"
    client._matches.clear()
    second = "https://www.orbea.com/en-be/orca-m30i-other"
    with patch.object(client, "_search_results", return_value=tuple(SearchProduct("U107TTCC", "Any name", url) for url in (URL, second))):
        assert client.lookup(SKU, "Orbea ORCA M30i").entry.product_link == URL


def test_collection_without_excel_saves_one_sku_package_with_all_assets(tmp_path):
    website, photos = Website(), Photos()
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan), patch.object(OrbeaAutomationService, "_download_images", tables):
        result = service(website, photos).run(config(tmp_path), resume=False)
    assert result.completed and result.counts["collected"] == 1
    assert website.lookups == [SKU] and website.closed
    folder = result.run_dir / "products" / SKU
    assert (folder / "description.html").read_text().strip() == "<p>A road bike for long days.</p>"
    assert (folder / "specifications.txt").read_text().strip() == "Frame: Orbea carbon"
    assert (folder / "photos/black-side.png").is_file()
    assert (folder / "tables/geometry-xs.png").is_file()
    metadata = json.loads((folder / "orbea-product.json").read_text())
    assert metadata["sku"] == SKU and metadata["status"] == "collected"
    assert metadata["pimbo_product_id"] == "p1"
    workbook = load_workbook(result.workbook_path)
    assert workbook["Collected Products"].cell(2, 5).value == "collected"
    assert workbook["Collected Products"].cell(2, 4).hyperlink
    workbook.close()


def test_confirmed_absent_description_is_saved_as_unavailable_instead_of_failed(tmp_path):
    website = Website(replace(Website().product, description_html="", description_source="not_available"))
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan):
        result = service(website, Photos()).run(config(tmp_path, download_images=False, download_product_photos=False), resume=False)
    assert result.completed
    metadata = json.loads((result.run_dir / "products" / SKU / "orbea-product.json").read_text())
    assert metadata["stages"]["description"]["status"] == "not_available"
    assert metadata["stages"]["description"]["source"] == "not_available"
    assert not metadata["errors"]


def test_retry_failed_replaces_old_incomplete_specification_capture(tmp_path):
    from tools.orbea_automation.website import SPECIFICATIONS_CAPTURE_VERSION
    settings = config(tmp_path, download_images=False, download_product_photos=False)
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan):
        first = service(Website(), Photos()).run(settings, resume=False)
    checkpoint = json.loads(first.checkpoint_path.read_text())
    checkpoint["results"][0]["collection_stages"]["specifications"]["capture_version"] = 1
    first.checkpoint_path.write_text(json.dumps(checkpoint))
    folder = first.run_dir / "products" / SKU
    (folder / "specifications.txt").write_text("Crankset: Incomplete option menu € 0 More information")
    metadata = json.loads((folder / "orbea-product.json").read_text())
    metadata["stages"]["specifications"]["capture_version"] = 1
    (folder / "orbea-product.json").write_text(json.dumps(metadata))
    source = (Path(__file__).parent / "fixtures/orbea_standard_configuration.html").read_text(encoding="utf-8")
    website = Website(parse_website_product(source, URL))
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")):
        resumed = service(website, Photos()).run(settings, resume=True, retry_failed=True)
    assert resumed.completed and resumed.run_dir == first.run_dir
    refreshed = (folder / "specifications.txt").read_text(encoding="utf-8")
    assert "Fork: RockShox ZEB" in refreshed and "Battery: Avinox 800 Wh" in refreshed
    assert "Incomplete option" not in refreshed
    metadata = json.loads((folder / "orbea-product.json").read_text())
    assert metadata["stages"]["specifications"]["capture_version"] == SPECIFICATIONS_CAPTURE_VERSION
    assert metadata["stages"]["specifications"]["source"] == "standard_configuration"


def test_failed_photos_resume_without_rescanning_or_rewriting_completed_text(tmp_path):
    website, photos = Website(), Photos()
    photos.fail = True
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan), patch.object(OrbeaAutomationService, "_download_images", tables):
        first = service(website, photos).run(config(tmp_path), resume=False)
    assert not first.completed
    text = first.run_dir / "products" / SKU / "description.html"
    text_time = text.stat().st_mtime_ns
    photos.fail = False
    website = Website()
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")), patch.object(OrbeaAutomationService, "_download_images", tables):
        resumed = service(website, photos).run(config(tmp_path), resume=True, retry_failed=True)
    assert resumed.completed and resumed.resumed and resumed.run_dir == first.run_dir
    assert not website.lookups and photos.calls == 2
    assert text.stat().st_mtime_ns == text_time


def test_retry_matched_refreshes_every_successful_package_without_lookup_or_scan(tmp_path):
    from tools.orbea_automation.checkpoint import RunCheckpoint

    def scan_two(client, catalogue, checkpoint, settings, **kwargs):
        scan(client, catalogue, checkpoint, settings, **kwargs)
        checkpoint.upsert_result(dict(row(), row_key="p2", product_id="p2"))

    settings = config(tmp_path, download_images=False)
    photos = Photos()
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan_two):
        first = service(Website(), photos).run(settings, resume=False)
    assert first.completed and first.counts["matched"] == 2 and photos.calls == 1
    saved = RunCheckpoint.load(first.run_dir, settings)
    saved.upsert_result(dict(row(), row_key="unmatched", product_id="unmatched", sku="T99901AA"))
    saved.mark_completed()
    old_folders = [first.run_dir / item["local_folder"] for item in saved.results if item["status"] == "code_match"]
    website = Website(replace(Website().product, description_html="<p>Updated description.</p>",
                              specifications_text="Frame: Updated frame"))
    events = []
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")):
        worker = service(website, photos)
        worker.pimbo_driver = None
        result = worker.run(settings, resume=True, retry_matched=True, progress=events.append)
    assert result.completed and result.run_dir == first.run_dir and result.counts["matched"] == 2
    assert website.lookups == [] and set(website.fetches) == {URL} and photos.calls == 2
    for folder in old_folders:
        assert (folder / "specifications.txt").read_text().strip() == "Frame: Updated frame"
        assert (folder / "description.html").read_text().strip() == "<p>Updated description.</p>"
        assert json.loads((folder / "orbea-product.json").read_text())["status"] == "collected"
    assert any(event.stage == "product_data" and event.current == event.total == 2 for event in events)
    assert json.loads(result.checkpoint_path.read_text())["results"][-1]["status"] == "unmatched"


def test_retry_matched_rechecks_unavailable_content_and_preserves_unselected_files(tmp_path):
    settings = config(tmp_path, download_images=False, download_product_photos=False)
    missing = Website(replace(Website().product, description_html="", description_source="not_available"))
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan):
        first = service(missing, Photos()).run(settings, resume=False)
    folder = first.run_dir / "products" / SKU
    existing_photo = folder / "photos/manual.png"
    existing_photo.parent.mkdir()
    existing_photo.write_bytes(b"keep this unselected photo")
    photo_time = existing_photo.stat().st_mtime_ns
    photos = Photos()
    website = Website()
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")):
        result = service(website, photos).run(settings, resume=True, retry_matched=True)
    assert result.completed and photos.calls == 0
    assert "A road bike" in (folder / "description.html").read_text()
    assert existing_photo.read_bytes() == b"keep this unselected photo"
    assert existing_photo.stat().st_mtime_ns == photo_time


def test_retry_matched_failed_refresh_blocks_old_specifications_and_keeps_them_on_disk(tmp_path):
    settings = config(tmp_path, download_images=False, download_product_photos=False)
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan):
        first = service(Website(), Photos()).run(settings, resume=False)
    website = Website(replace(Website().product, specifications_error="Full configuration did not load"))
    result = service(website, Photos()).run(settings, resume=True, retry_matched=True)
    assert not result.completed
    folder = first.run_dir / "products" / SKU
    assert (folder / "specifications.txt").read_text().strip() == "Frame: Orbea carbon"
    metadata = json.loads((folder / "orbea-product.json").read_text())
    assert metadata["stages"]["specifications"]["status"] == "error"
    assert metadata["stages"]["specifications"]["files"] == []


def test_stopped_matched_retry_leaves_all_packages_pending_for_resume(tmp_path):
    from tools.orbea_automation.models import RunCancelled
    settings = config(tmp_path, download_images=False, download_product_photos=False)

    def scan_two(client, catalogue, checkpoint, run_config, **kwargs):
        scan(client, catalogue, checkpoint, run_config, **kwargs)
        checkpoint.upsert_result(dict(row(), row_key="p2", product_id="p2"))

    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan_two):
        first = service(Website(), Photos()).run(settings, resume=False)
    website = Website()
    with patch.object(website, "fetch", side_effect=RunCancelled("stopped")):
        stopped = service(website, Photos()).run(settings, resume=True, retry_matched=True)
    assert stopped.cancelled
    data = json.loads(stopped.checkpoint_path.read_text())
    assert all(item["collection_stages"]["specifications"]["status"] == "pending" for item in data["results"])
    for item in data["results"]:
        metadata = json.loads((stopped.run_dir / item["local_folder"] / "orbea-product.json").read_text())
        assert metadata["status"] == "partial" and metadata["stages"]["specifications"]["status"] == "pending"
    website = Website(replace(Website().product, specifications_text="Frame: Refreshed after resume"))
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")):
        result = service(website, Photos()).run(settings, resume=True)
    assert result.completed and result.counts["collected"] == 2
    data = json.loads(result.checkpoint_path.read_text())
    for item in data["results"]:
        assert (result.run_dir / item["local_folder"] / "specifications.txt").read_text().strip() == "Frame: Refreshed after resume"


def test_retry_matched_recaptures_completed_tables_once_per_url(tmp_path):
    from tools.orbea_automation.checkpoint import RunCheckpoint, create_run_directory
    from tools.orbea_table_image_downloader import AVAILABILITY_PROBE_VERSION, GEOMETRY_CAPTURE_VERSION
    settings = config(tmp_path, download_product_photos=False, download_description=False, download_specifications=False)
    saved = RunCheckpoint.create(create_run_directory(settings.output_root), settings)
    for key in ("p1", "p2"):
        saved.upsert_result(dict(row(), row_key=key, product_id=key, status="code_match", catalogue_url=URL,
            collection_stages={"tables": {"status": "not_available", "files": []}}))
    saved.upsert_image(URL, {"geometry_status": "not_available", "size_guide_status": "downloaded",
        "availability_probe_version": AVAILABILITY_PROBE_VERSION, "geometry_capture_version": GEOMETRY_CAPTURE_VERSION,
        "attempts": 5})
    saved.data.update(scan_completed=True, website_lookup_completed=True)
    saved.mark_completed()
    captured = []
    driver = SimpleNamespace(quit=lambda: None)

    def capture(browser, url, geometry_path, size_path, **kwargs):
        assert browser is driver and kwargs["need_geometry"] and kwargs["need_size_guide"]
        captured.append(url)
        geometry_path.write_bytes(b"refreshed geometry")
        size_path.write_bytes(b"refreshed sizes")
        return {"geometry_status": "downloaded", "size_guide_status": "downloaded", "retryable": False, "errors": []}

    with patch("tools.orbea_table_image_downloader.capture_orbea_tables", capture):
        worker = OrbeaAutomationService(None, image_driver_factory=lambda *_: driver, website_client_factory=lambda *_: Website())
        result = worker.run(settings, resume=True, retry_matched=True)
    assert result.completed and captured == [URL]
    data = json.loads(result.checkpoint_path.read_text())
    assert data["images"][URL]["attempts"] == 1
    for item in data["results"]:
        folder = result.run_dir / item["local_folder"] / "tables"
        assert (folder / "geometry.png").read_bytes() == b"refreshed geometry"
        assert (folder / "size-guide-cm.png").read_bytes() == b"refreshed sizes"


def test_retry_matched_without_a_saved_match_does_not_create_a_fresh_run(tmp_path):
    settings = config(tmp_path)
    with pytest.raises(ValueError, match="No saved matched"):
        service(Website(), Photos()).run(settings, resume=True, retry_matched=True)
    assert not settings.output_root.exists()


def test_retry_matched_uses_saved_matches_even_when_other_lookups_or_scan_are_incomplete(tmp_path):
    from tools.orbea_automation.checkpoint import RunCheckpoint, create_run_directory, saved_run_config
    settings = config(tmp_path, download_images=False, download_product_photos=False)
    saved = RunCheckpoint.create(create_run_directory(settings.output_root), settings)
    saved.upsert_result(dict(row(), status="code_match", catalogue_url=URL, catalogue_model="ORCA M30i"))
    saved.upsert_result(dict(row(), row_key="p2", product_id="p2", sku="T99901AA", website_lookup_status="error"))
    website = Website()
    worker = service(website, Photos())
    worker.pimbo_driver = None
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")):
        result = worker.run(saved_run_config(saved.run_dir), resume=True, retry_matched=True)
    assert result.counts["collected"] == 1 and website.fetches and not website.lookups
    assert not result.completed  # The original scan/other lookup still remain unfinished.
    data = json.loads(result.checkpoint_path.read_text())
    assert data["results"][0]["collection_status"] == "collected"
    assert data["results"][1]["website_lookup_status"] == "error"
    assert (result.run_dir / "products" / SKU / "specifications.txt").read_text().strip() == "Frame: Orbea carbon"


def test_unconfirmed_product_is_kept_for_review_and_not_downloaded(tmp_path):
    website, photos = Website(parse_website_product(page("T12302AA"), URL)), Photos()
    def no_tables(_service, config, checkpoint, *args, **kwargs):
        assert _service._image_jobs(checkpoint) == []
        checkpoint.data["images_completed"] = True
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan), patch.object(OrbeaAutomationService, "_download_images", no_tables):
        result = service(website, photos).run(config(tmp_path), resume=False)
    assert not result.completed and photos.calls == 0
    data = json.loads(result.checkpoint_path.read_text())
    assert data["results"][0]["status"] == "ambiguous"
    assert not (result.run_dir / "products").exists()


def test_stop_during_photo_download_preserves_completed_stages(tmp_path):
    website, photos = Website(), Photos()
    photos.cancel_run = True
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan), patch.object(OrbeaAutomationService, "_download_images", tables):
        stopped = service(website, photos).run(config(tmp_path), resume=False)
    assert stopped.cancelled and not stopped.completed and website.closed
    data = json.loads(stopped.checkpoint_path.read_text())
    assert data["results"][0]["collection_stages"]["description"]["status"] == "downloaded"
    assert stopped.workbook_path.is_file()
    photos.cancel_run = False
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")), patch.object(OrbeaAutomationService, "_download_images", tables):
        resumed = service(Website(), photos).run(config(tmp_path), resume=True)
    assert resumed.completed and not resumed.cancelled and resumed.run_dir == stopped.run_dir


def test_unselected_assets_are_not_downloaded(tmp_path):
    website, photos = Website(), Photos()
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan), patch.object(OrbeaAutomationService, "_download_images", side_effect=AssertionError("tables disabled")):
        result = service(website, photos).run(config(tmp_path, download_images=False, download_product_photos=False, download_description=False, download_specifications=False), resume=False)
    assert result.completed and photos.calls == 0
    folder = result.run_dir / "products" / SKU
    assert list(folder.iterdir()) == [folder / "orbea-product.json"]


def test_access_failure_can_be_retried_without_losing_the_scanned_code(tmp_path):
    website, photos = Website(), Photos()
    with patch.object(website, "lookup", side_effect=OrbeaAccessError("Orbea access check")), patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan):
        with pytest.raises(OrbeaRunFailure, match="access check") as failure:
            service(website, photos).run(config(tmp_path, download_images=False), resume=False)
    first = failure.value.partial_result
    assert not first.completed and first.counts["errors"] == 1
    assert photos.calls == 0 and first.workbook_path.is_file() and website.closed
    assert isinstance(failure.value.__cause__, OrbeaAccessError)
    assert failure.value.reason == "website_access"
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")):
        resumed = service(Website(), photos).run(config(tmp_path, download_images=False), resume=True, retry_failed=True)
    assert resumed.completed and resumed.run_dir == first.run_dir
    data = json.loads(resumed.checkpoint_path.read_text())
    assert "last_error" not in data and "Orbea access check" in data["previous_error"]


def test_explicit_saved_scan_resumes_without_pimbo_or_the_original_catalogue(tmp_path):
    from tools.orbea_automation.checkpoint import RunCheckpoint, create_run_directory, saved_run_config
    catalogue = tmp_path / "catalogue.xlsx"
    catalogue.write_bytes(b"catalogue fixture; the completed scan already contains its mappings")
    run_config = config(tmp_path, download_images=False, download_product_photos=False,
        download_description=False, download_specifications=False)
    run_config = replace(run_config, catalogue_path=catalogue)
    checkpoint = RunCheckpoint.create(create_run_directory(run_config.output_root), run_config)
    checkpoint.upsert_result(row())
    checkpoint.data["scan_completed"] = True
    checkpoint.save()
    catalogue.unlink()
    restored = saved_run_config(checkpoint.run_dir)
    website = Website()
    service = OrbeaAutomationService(None, website_client_factory=lambda *_: website)
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")), patch("tools.orbea_automation.service.CatalogueIndex.from_workbook", side_effect=AssertionError("must use saved scan")):
        result = service.run(restored, resume=True)
    assert result.completed and result.resumed and result.run_dir == checkpoint.run_dir
    assert result.workbook_path.is_file() and website.fetches == []


def test_access_failure_stays_partial_even_when_all_assets_are_unselected(tmp_path):
    website, photos = Website(), Photos()
    with patch.object(website, "lookup", side_effect=OrbeaAccessError("Orbea access check")), patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan):
        with pytest.raises(OrbeaRunFailure) as failure:
            service(website, photos).run(config(tmp_path, download_images=False, download_product_photos=False, download_description=False, download_specifications=False), resume=False)
    result = failure.value.partial_result
    assert not result.completed and result.counts["collected"] == 0
    data = json.loads(result.checkpoint_path.read_text())
    assert data["results"][0]["collection_status"] == "partial"


def test_security_check_stops_whole_website_stage_and_preserves_other_rows(tmp_path):
    website, photos = Website(), Photos()

    def two_product_scan(_client, catalogue, checkpoint, config, **kwargs):
        scan(_client, catalogue, checkpoint, config, **kwargs)
        second = dict(row(), row_key="p2", product_id="p2", sku="U90709FJ")
        checkpoint.upsert_result(second)

    with patch.object(website, "lookup", side_effect=OrbeaAccessError("Orbea access check")) as lookup, patch("tools.orbea_automation.service.PimboBrowserClient.collect", two_product_scan), patch.object(OrbeaAutomationService, "_download_images") as downloads:
        with pytest.raises(OrbeaRunFailure) as failure:
            service(website, photos).run(config(tmp_path), resume=False)
    assert lookup.call_count == 1 and not downloads.called and photos.calls == 0
    result = failure.value.partial_result
    data = json.loads(result.checkpoint_path.read_text())
    assert data["scan_completed"] and not data["website_lookup_completed"]
    assert len(data["results"]) == 2 and "website_lookup_status" not in data["results"][1]
    workbook = load_workbook(result.workbook_path)
    assert workbook["Raw Scan"].max_row == 3
    workbook.close()


def test_missing_downloaded_table_is_recaptured_on_retry(tmp_path):
    from tools.orbea_automation.checkpoint import RunCheckpoint, create_run_directory
    from tools.orbea_table_image_downloader import AVAILABILITY_PROBE_VERSION, GEOMETRY_CAPTURE_VERSION
    run_config = config(tmp_path)
    checkpoint = RunCheckpoint.create(create_run_directory(run_config.output_root), run_config)
    record = row()
    record.update(status="code_match", catalogue_url=URL, catalogue_model="ORCA M30i")
    checkpoint.upsert_result(record)
    checkpoint.upsert_image(URL, {"geometry_status": "downloaded", "size_guide_status": "not_available", "attempts": 2,
        "availability_probe_version": AVAILABILITY_PROBE_VERSION, "geometry_capture_version": GEOMETRY_CAPTURE_VERSION})
    driver = SimpleNamespace(quit=lambda: None)
    reporter = SimpleNamespace(emit=lambda *args: None)
    def capture(_driver, url, geometry_path, size_path, **kwargs):
        assert kwargs["need_geometry"] and not kwargs["need_size_guide"]
        geometry_path.write_bytes(b"recaptured table")
        return {"geometry_status": "downloaded", "size_guide_status": "not_available", "retryable": False, "errors": []}
    worker = OrbeaAutomationService(object(), image_driver_factory=lambda *_: driver)
    with patch("tools.orbea_table_image_downloader.capture_orbea_tables", capture) as _:
        worker._download_images(run_config, checkpoint, reporter, CancellationToken(), None, retry_failed=True)
    assert checkpoint.images[URL]["attempts"] == 1
    assert (checkpoint.run_dir / checkpoint.images[URL]["geometry_image"]).is_file()


def test_all_urls_are_gathered_before_any_product_download(tmp_path):
    class BlockedSecondProduct(Website):
        def lookup(self, sku, title):
            if sku == "U90709FJ":
                assert not list((tmp_path / "runs").glob("*/products/*"))
                assert not self.fetches
                raise OrbeaAccessError("Second product blocked")
            return super().lookup(sku, title)

    def two_product_scan(_client, catalogue, checkpoint, config, **kwargs):
        scan(_client, catalogue, checkpoint, config, **kwargs)
        checkpoint.upsert_result(dict(row(), row_key="p2", product_id="p2", sku="U90709FJ"))

    website = BlockedSecondProduct()
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", two_product_scan):
        with pytest.raises(OrbeaRunFailure, match="Second product blocked") as failure:
            service(website, Photos()).run(config(tmp_path, download_images=False, download_product_photos=False), resume=False)
    assert failure.value.partial_result.counts["collected"] == 0
    assert failure.value.partial_result.counts["scanned"] == 2


def test_description_dialog_failure_does_not_block_other_assets_and_can_resume(tmp_path):
    website, photos = Website(), Photos()
    website.product = replace(website.product, description_html="", description_error="Features dialog: timed out")
    run_config = config(tmp_path, download_images=False)
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan):
        first = service(website, photos).run(run_config, resume=False)
    assert not first.completed and photos.calls == 1
    data = json.loads(first.checkpoint_path.read_text())
    stages = data["results"][0]["collection_stages"]
    assert stages["description"]["status"] == "error"
    assert stages["specifications"]["status"] == stages["photos"]["status"] == "downloaded"
    assert data["website_lookup_completed"]
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")):
        resumed = service(Website(), photos).run(run_config, resume=True, retry_failed=True)
    assert resumed.completed and resumed.run_dir == first.run_dir and photos.calls == 1


def test_retry_refreshes_legacy_description_in_a_completed_package(tmp_path):
    from tools.orbea_automation.features import DESCRIPTION_CAPTURE_VERSION

    run_config = config(tmp_path, download_images=False, download_product_photos=False)
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", scan):
        first = service(Website(), Photos()).run(run_config, resume=False)
    assert first.completed
    data = json.loads(first.checkpoint_path.read_text())
    data["results"][0]["collection_stages"]["description"].pop("capture_version")
    first.checkpoint_path.write_text(json.dumps(data), encoding="utf-8")
    description_file = first.run_dir / "products" / SKU / "description.html"
    description_file.write_text("OLD SHORT SUMMARY", encoding="utf-8")
    website = Website(replace(Website().product, description_html="<h2>Full Features</h2><p>All feature cards.</p>", description_source="feature_dialog"))
    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", side_effect=AssertionError("must not rescan")):
        resumed = service(website, Photos()).run(run_config, resume=True, retry_failed=True)
    assert resumed.completed and resumed.resumed and resumed.run_dir == first.run_dir
    assert "Full Features" in description_file.read_text()
    data = json.loads(resumed.checkpoint_path.read_text())
    assert data["results"][0]["collection_stages"]["description"]["capture_version"] == DESCRIPTION_CAPTURE_VERSION
    assert data["results"][0]["collection_stages"]["description"]["source"] == "feature_dialog"


def test_product_table_capture_filters_jobs_and_keeps_borrowed_browser_open(tmp_path):
    from tools.orbea_automation.checkpoint import RunCheckpoint, create_run_directory

    run_config = config(tmp_path)
    checkpoint = RunCheckpoint.create(create_run_directory(run_config.output_root), run_config)
    checkpoint.upsert_result(dict(row(), status="code_match", catalogue_url=URL, catalogue_model="ORCA M30i"))
    other = "https://www.orbea.com/en-be/alma-carbon"
    checkpoint.upsert_result(dict(row(), row_key="p2", product_id="p2", sku="U90709FJ", status="code_match", catalogue_url=other, catalogue_model="ALMA CARBON"))
    from unittest.mock import Mock
    driver = Mock()
    captured = []

    def capture(browser, url, geometry_path, size_path, **kwargs):
        assert browser is driver
        captured.append(url)
        geometry_path.write_bytes(b"geometry fixture")
        size_path.write_bytes(b"size fixture")
        return {"geometry_status": "downloaded", "size_guide_status": "downloaded", "retryable": False, "errors": []}

    worker = OrbeaAutomationService(object())
    with patch("tools.orbea_table_image_downloader.capture_orbea_tables", capture):
        worker._download_images(run_config, checkpoint, SimpleNamespace(emit=lambda *args: None), CancellationToken(), None,
            retry_failed=True, row_keys={"p1"}, driver_factory=lambda: driver)
    assert captured == [URL] and other not in checkpoint.images
    driver.quit.assert_not_called()


def test_scan_failure_preserves_excel_and_emits_partial_result_before_error(tmp_path):
    from GUI_Qt.orbea.workers import OrbeaRunWorker

    def failing_scan(_client, catalogue, checkpoint, config, **kwargs):
        checkpoint.upsert_result(row())
        raise IndexError("Pimbo product row 20 disappeared")

    with patch("tools.orbea_automation.service.PimboBrowserClient.collect", failing_scan):
        with pytest.raises(OrbeaRunFailure, match="row 20 disappeared") as failure:
            service(Website(), Photos()).run(config(tmp_path), resume=False)
    partial = failure.value.partial_result
    assert partial.workbook_path.is_file() and partial.checkpoint_path.is_file()
    assert partial.counts["scanned"] == 1 and not partial.completed
    assert isinstance(failure.value.__cause__, IndexError)

    events = []
    failed_service = SimpleNamespace(run=lambda *args, **kwargs: (_ for _ in ()).throw(failure.value))
    worker = OrbeaRunWorker(object(), lambda *_: failed_service, config(tmp_path), resume=False, retry_failed=False)
    worker.partial_result.connect(lambda result: events.append(("partial", result)))
    worker.failed.connect(lambda message: events.append(("failed", message)))
    worker.succeeded.connect(lambda result: events.append(("succeeded", result)))
    worker.run()
    assert [event[0] for event in events] == ["partial", "failed"]
    assert events[0][1].workbook_path == partial.workbook_path


def test_successful_url_stage_finishes_before_first_fetch(tmp_path):
    second_url = 'https://www.orbea.com/en-au/alma-carbon'
    class OrderedWebsite(Website):
        def lookup(self, sku, title):
            self.lookups.append(sku)
            code, name, url = ('U107TTCC', 'ORCA M30i', URL) if sku == SKU else ('U907TTCC', 'Alma Carbon', second_url)
            return MatchResult('code_match', 'Orbea TTCC search', CatalogueEntry(code[:4], code, name, product_link=url))

        def fetch(self, url):
            assert self.lookups == [SKU, 'U90709FJ']
            self.fetches.append(url)
            return self.product if url == URL else parse_website_product(page('U907TTCC', 'Alma Carbon'), second_url)

    def two_product_scan(_client, catalogue, checkpoint, config, **kwargs):
        scan(_client, catalogue, checkpoint, config, **kwargs)
        checkpoint.upsert_result(dict(row(), row_key='p2', product_id='p2', sku='U90709FJ'))

    website = OrderedWebsite()
    with patch('tools.orbea_automation.service.PimboBrowserClient.collect', two_product_scan):
        result = service(website, Photos()).run(config(tmp_path, download_images=False, download_product_photos=False), resume=False)
    assert result.completed and result.counts['collected'] == 2
    data = json.loads(result.checkpoint_path.read_text())
    assert [row['catalogue_url'] for row in data['results']] == [URL, second_url]
    assert all(row['website_lookup_status'] == 'done' for row in data['results'])


def test_incomplete_url_stage_does_not_download_old_catalogue_candidates(tmp_path):
    website, photos = Website(), Photos()
    def legacy_scan(_client, catalogue, checkpoint, config, **kwargs):
        scan(_client, catalogue, checkpoint, config, **kwargs)
        checkpoint.results[0].update(status='code_match', catalogue_url=URL, website_lookup_status='done')
    with patch.object(website, 'lookup', side_effect=RuntimeError('Regional lookup is incomplete')), patch('tools.orbea_automation.service.PimboBrowserClient.collect', legacy_scan), patch.object(OrbeaAutomationService, '_download_images') as downloads:
        result = service(website, photos).run(config(tmp_path), resume=False)
    assert not result.completed
    assert website.fetches == [] and photos.calls == 0 and not downloads.called
    data = json.loads(result.checkpoint_path.read_text())
    assert not data['website_lookup_completed']
    assert data['results'][0]['website_lookup_status'] == 'error'


def test_resume_reuses_verified_link_for_another_variant_of_same_model(tmp_path):
    from tools.orbea_automation.checkpoint import RunCheckpoint, create_run_directory
    from tools.orbea_automation.website import WEBSITE_LOOKUP_VERSION
    run_config = config(tmp_path)
    checkpoint = RunCheckpoint.create(create_run_directory(run_config.output_root), run_config)
    first = dict(row(), status='code_match', catalogue_code='U107TTCC', catalogue_model='ORCA M30i', catalogue_url=URL,
                 website_lookup_status='done', website_lookup_version=WEBSITE_LOOKUP_VERSION)
    checkpoint.upsert_result(first)
    checkpoint.upsert_result(dict(row(), row_key='p2', product_id='p2', sku='U10708AA'))
    website = Website()
    worker = service(website, Photos())
    worker._lookup_products(website, checkpoint, SimpleNamespace(emit=lambda *_: None), CancellationToken(), None)
    assert website.lookups == [] and website.fetches == []
    assert all(row['catalogue_url'] == URL for row in checkpoint.results)
