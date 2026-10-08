from __future__ import annotations

import json
from types import SimpleNamespace

from openpyxl import load_workbook

from Utilities.UploadResultsReport import history_upload_row, upload_result_row, write_upload_results, read_supplier_result


def match(sku="SKU", folder=""):
    return SimpleNamespace(sku=sku, ready=True, local_folder=folder, pimbo_product_name="Orbea bike",
                           pimbo_product_id="p1", pimbo_product_url="https://pim.bo.ultrabike.lt/dashboard/products/p1", note="")


def test_excel_export_includes_success_failed_skipped_and_unprocessed_and_review_list(tmp_path):
    records = [
        {"preparation": {"status": "saved_automatically", "photo_upload": {"action": "uploaded", "placeholders_removed": 1, "uploaded_photos": 3}},
         "selected_stages": ["product_photos", "save"], "completed_stages": ["product_photos", "save"]},
        {"preparation": {"status": "no_changes", "photo_upload": {"action": "skipped_existing", "existing_photos": 2}},
         "selected_stages": ["product_photos", "save"], "completed_stages": ["product_photos"]},
        {"preparation": {"status": "failed", "error": "Could not remove placeholder", "photo_upload": {"action": "failed"}},
         "selected_stages": ["product_photos", "save"], "completed_stages": []},
        None,
    ]
    rows = [upload_result_row(match(str(index)), record, brand="Orbea") for index, record in enumerate(records)]
    path = tmp_path / "results.xlsx"
    write_upload_results(path, rows, context={"Batch stop reason": "Product could not be opened"})
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["Summary", "Results", "Needs checking"]
    headers = [cell.value for cell in workbook["Results"][1]]
    results = [dict(zip(headers, row)) for row in workbook["Results"].iter_rows(min_row=2, values_only=True)]
    assert [row["Needs checking"] for row in results] == ["No", "No", "Yes", "Yes"]
    assert results[0]["Placeholders removed"] == 1 and results[0]["Photos added to form"] == 3
    assert results[0]["Photo changes saved"] == "Yes"
    assert results[1]["Photo result"] == "Skipped: real photos already present"
    assert results[1]["Photo changes saved"] == "Not needed"
    assert "Could not remove placeholder" in results[2]["Reason to check"]
    assert workbook["Needs checking"].max_row == 3
    assert workbook["Results"].freeze_panes == "A2"
    assert workbook["Results"].auto_filter.ref
    assert dict(workbook["Summary"].values)["Batch stop reason"] == "Product could not be opened"
    workbook.close()


def test_failed_second_phase_keeps_saved_photo_success_but_still_needs_check():
    row = upload_result_row(match(), {"preparation": {"status": "failed", "error": "Specifications failed",
        "photo_upload": {"action": "uploaded", "uploaded_photos": 2}},
        "selected_stages": ["product_photos", "specifications_magic_ai", "save"],
        "completed_stages": ["product_photos", "save"]})
    assert row["Needs checking"] == "Yes" and row["Photo changes saved"] == "Yes"
    assert "specifications_magic_ai" in row["Reason to check"]


def test_unsaved_placeholder_removal_needs_review():
    row = upload_result_row(match(), {"preparation": {"status": "ready_for_review", "photo_upload": {
        "action": "skipped_existing", "placeholders_removed": 1}}})
    assert row["Needs checking"] == "Yes"
    assert row["Photo changes saved"] == "No — review required"


def test_exported_text_is_never_an_excel_formula(tmp_path):
    path = tmp_path / "literal.xlsx"
    row = upload_result_row(match(), {"preparation": {"status": "failed", "error": '=HYPERLINK("https://example.com")'}})
    write_upload_results(path, [row], context={"Batch stop reason": "=1+1"})
    workbook = load_workbook(path, data_only=False)
    assert all(cell.data_type != "f" for sheet in workbook for cells in sheet for cell in cells)
    workbook.close()


def test_generic_history_export_reads_photo_details_and_legacy_results():
    base = {"brand": "KROSS", "product_code": "SKU", "status": "saved_manually", "processed_at": "2026-10-07"}
    record = {**base, "details_json": json.dumps({"pim_preparation": {"status": "saved_manually", "photo_upload": {
        "action": "uploaded", "placeholders_removed": 1, "uploaded_photos": 4}, "final_url": "https://pim.bo.ultrabike.lt/dashboard/products/p1"}})}
    row = history_upload_row(record)
    assert row["Photo changes saved"] == "Yes" and row["Photos added to form"] == 4
    assert row["Needs checking"] == "No" and row["Pimbo URL"].endswith("/p1")
    assert history_upload_row({**base, "details_json": "invalid"})["Photo result"] == "Not recorded"
    assert history_upload_row({**base, "error_message": "Error"})["Needs checking"] == "Yes"


def test_saved_supplier_report_rejects_results_for_another_sku_or_product(tmp_path):
    path = tmp_path / "kross-upload-result.json"
    target = match(folder=str(tmp_path))
    path.write_text(json.dumps({"preparation": {"product_code": "SKU", "product_id": "p1", "status": "saved_automatically"}}))
    assert read_supplier_result(target, "kross")
    assert read_supplier_result(match("OTHER", str(tmp_path)), "kross") is None
    target.pimbo_product_id = "p2"
    assert read_supplier_result(target, "kross") is None
