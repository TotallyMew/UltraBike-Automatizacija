from __future__ import annotations

import json
import os
import time
from types import SimpleNamespace

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PySide6.QtWidgets import QApplication, QWidget
from GUI_Qt.screens.OrbeaScreen import OrbeaScreen
from GUI_Qt.orbea.upload import OrbeaUploadWorker
from Managers.PimboProductEditor import PimPreparationResult, PimPreparationStatus
from tools.orbea_automation.upload import OrbeaUploadService, OrbeaUploadResult, OrbeaWorkflowOptions


def make_package(root, sku="U10707SV"):
    folder = root / "products" / sku
    folder.mkdir(parents=True)
    (folder / "orbea-product.json").write_text(json.dumps({"sku": sku, "pimbo_product_id": "p1", "pimbo_url": "https://pim.bo.ultrabike.lt/dashboard/products/p1",
        "orbea_url": "https://www.orbea.com/en-be/orca-m30i", "orbea_product_name": "ORCA M30i", "status": "collected", "stages": {}}))
    return folder


class Settings:
    def __init__(self):
        self.values = {}

    def get(self, key, default=None):
        return self.values.get(key, default)

    def set(self, key, value):
        self.values[key] = value


@pytest.fixture
def screen(tmp_path):
    app = QApplication.instance() or QApplication([])
    main = QWidget()
    main.settings = Settings()
    main.i18n = SimpleNamespace(tr=lambda key, **kwargs: key)
    main.driver = None
    main.leases = []
    main.try_acquire_browser_lease = lambda owner: main.leases.append("acquire") or True
    main.release_browser_lease = lambda owner: main.leases.append("release")
    widget = OrbeaScreen(main)
    app.processEvents()
    yield widget
    assert widget.shutdown()
    widget.deleteLater()
    main.deleteLater()
    app.processEvents()


def test_current_collection_loads_products_and_has_all_workflow_steps(screen, tmp_path):
    make_package(tmp_path)
    screen._load_run_result(SimpleNamespace(run_dir=tmp_path, counts={}))
    panel = screen._upload_panel
    assert panel.matches[0].sku == "U10707SV" and len(panel.selected_matches()) == 1
    assert panel.options().selected_stages == OrbeaWorkflowOptions.STAGES
    assert not panel.start.isEnabled()
    screen.main.driver = object()
    screen._update_action_states()
    assert panel.start.isEnabled()
    screen._switch_section("upload")
    assert not screen._automation_page.isHidden() and screen._tools_page.isHidden()


def test_workflow_choices_persist_and_unsaved_multi_product_run_is_blocked(screen, tmp_path):
    make_package(tmp_path, "U10707SV")
    make_package(tmp_path, "U10707SX")
    panel = screen._upload_panel
    panel.load_packages(tmp_path)
    panel._select_steps(False)
    panel.checks["brand"].setChecked(True)
    assert panel.options().selected_stages == ("brand",)
    assert json.loads(screen.main.settings.values["orbea_workflow_options"])["brand"] is True
    screen.main.driver = object()
    messages = []
    screen._warn = lambda *args: messages.append(args)
    panel.start_upload()
    assert messages and "only one product" in messages[0][1]
    assert not panel.worker and not screen.main.leases


def test_upload_worker_releases_browser_and_disables_collection_while_running(screen, tmp_path):
    make_package(tmp_path)
    panel = screen._upload_panel
    panel.load_packages(tmp_path)
    panel._select_steps(False)
    panel.checks["brand"].setChecked(True)
    screen.main.driver = object()
    calls = []
    class Service:
        def upload_and_save(self, match, root, **kwargs):
            calls.append(kwargs["options"].selected_stages)
            time.sleep(0.03)
            return OrbeaUploadResult(match, PimPreparationResult(match.sku, status=PimPreparationStatus.READY_FOR_REVIEW), options=kwargs["options"])
    panel.service_factory = Service
    panel.start_upload()
    assert screen.is_running() and not screen._start_btn.isEnabled()
    assert not panel.table.isEnabled() and panel.stop.isEnabled()
    deadline = time.monotonic() + 3
    while panel.is_running() and time.monotonic() < deadline:
        QApplication.processEvents()
        time.sleep(0.005)
    assert not panel.is_running()
    assert calls == [("brand",)] and screen.main.leases == ["acquire", "release"]
    assert panel.results[0].succeeded and not panel.selected_matches()
    assert "1 successful" in panel.status.text()


def test_worker_stop_is_between_products_and_emits_cancelled_summary(tmp_path):
    make_package(tmp_path, "U10707SV")
    make_package(tmp_path, "U10707SX")
    matches = OrbeaUploadService.load_local_packages(tmp_path)
    calls, summaries = [], []
    class Service:
        def upload_and_save(self, match, root, **kwargs):
            calls.append(match.sku)
            return OrbeaUploadResult(match, PimPreparationResult(match.sku, status=PimPreparationStatus.SAVED_AUTOMATICALLY))
    worker = OrbeaUploadWorker(Service, matches, None)
    worker.item_finished.connect(lambda result: worker.request_stop())
    worker.completed.connect(summaries.append)
    worker.run()
    assert calls == [matches[0].sku]
    assert summaries[0]["status"] == "cancelled" and summaries[0]["unprocessed"] == 1


@pytest.mark.parametrize("status,warnings,expected", [
    (PimPreparationStatus.SAVED_AUTOMATICALLY, (), "succeeded"),
    (PimPreparationStatus.SAVED_AUTOMATICALLY, ("Missing size chart",), "partial"),
    (PimPreparationStatus.FAILED, (), "failed"),
])
def test_worker_summary_reports_product_failure_and_warnings(tmp_path, status, warnings, expected):
    make_package(tmp_path)
    match = OrbeaUploadService.load_local_packages(tmp_path)[0]
    result = OrbeaUploadResult(match, PimPreparationResult(match.sku, status=status, warnings=warnings))
    service = SimpleNamespace(upload_and_save=lambda *args, **kwargs: result)
    worker = OrbeaUploadWorker(lambda: service, (match,), None)
    summaries = []
    worker.completed.connect(summaries.append)
    worker.run()
    assert summaries[0]["status"] == expected


def test_old_specifications_block_the_whole_batch_before_browser_acquisition(screen, tmp_path):
    make_package(tmp_path, "U10707SV")
    make_package(tmp_path, "U10707SX")
    panel = screen._upload_panel
    panel.load_packages(tmp_path)
    screen.main.driver = object()
    messages = []
    screen._warn = lambda *args: messages.append(args)
    panel.start_upload()
    assert not panel.worker and not screen.main.leases
    assert messages and "2 selected product(s)" in messages[0][1]
    assert "Retry matched downloads" in panel.status.text()


def test_finished_status_shows_the_actual_failure_reason(screen, tmp_path):
    make_package(tmp_path)
    panel = screen._upload_panel
    panel.load_packages(tmp_path)
    match = panel.matches[0]
    panel.results = [OrbeaUploadResult(match, PimPreparationResult(match.sku,
        status=PimPreparationStatus.FAILED, error="Source specifications did not load"))]
    panel.worker = SimpleNamespace(matches=(match,), _batch_error="", isRunning=lambda: False)
    panel._finished()
    assert "1 failed or blocked" in panel.status.text()
    assert "Reason: Source specifications did not load" in panel.status.text()


def write_upload_record(folder, sku, status, completed):
    record = {"preparation": {"product_code": sku, "product_id": "p1", "status": status, "error": "Spec AI failed" if status == "failed" else ""},
              "selected_stages": list(OrbeaWorkflowOptions.STAGES), "completed_stages": list(completed)}
    (folder / "orbea-upload-result.json").write_text(json.dumps(record))


def test_loading_saved_products_keeps_successful_uploads_unselected(screen, tmp_path):
    saved = make_package(tmp_path, "U10707SV")
    failed = make_package(tmp_path, "U10707SX")
    make_package(tmp_path, "U10707SY")
    write_upload_record(saved, "U10707SV", "saved_automatically", OrbeaWorkflowOptions.STAGES)
    write_upload_record(failed, "U10707SX", "failed", ("description_source", "save"))
    panel = screen._upload_panel
    panel.load_packages(tmp_path)
    assert {match.sku for match in panel.selected_matches()} == {"U10707SX", "U10707SY"}
    assert "saved automatically" in panel.table.item(0, 3).text().lower()
    assert "Spec AI failed" in panel.table.item(1, 4).text()
    panel.load_packages(tmp_path)
    assert {match.sku for match in panel.selected_matches()} == {"U10707SX", "U10707SY"}


def test_retry_button_runs_only_failed_and_unprocessed_products_with_saved_step_resume(screen, tmp_path):
    saved = make_package(tmp_path, "U10707SV")
    failed = make_package(tmp_path, "U10707SX")
    make_package(tmp_path, "U10707SY")
    write_upload_record(saved, "U10707SV", "saved_automatically", OrbeaWorkflowOptions.STAGES)
    write_upload_record(failed, "U10707SX", "failed", ("description_source", "save"))
    panel = screen._upload_panel
    panel.load_packages(tmp_path)
    panel._select_steps(False)
    panel.checks["brand"].setChecked(True)
    panel.checks["save"].setChecked(True)
    screen.main.driver = object()
    calls = []
    class Service:
        def upload_and_save(self, match, root, **kwargs):
            calls.append((match.sku, kwargs.get("resume_completed")))
            return OrbeaUploadResult(match, PimPreparationResult(match.sku, status=PimPreparationStatus.SAVED_AUTOMATICALLY),
                                     options=kwargs["options"], completed_stages=kwargs["options"].selected_stages)
    panel.service_factory = Service
    panel._select(True)  # The retry action still excludes a successful selected bike.
    assert panel.retry_upload.isEnabled()
    panel.retry_upload.click()
    assert not panel.retry_upload.isEnabled()
    assert panel.worker.wait(5000)
    QApplication.processEvents()
    assert calls == [("U10707SX", True), ("U10707SY", True)]
    assert len(panel.worker.matches) == 2
    assert "2 successful" in panel.status.text()


def test_retry_worker_excludes_saved_successes_even_when_called_with_every_match(tmp_path):
    saved = make_package(tmp_path, "U10707SV")
    make_package(tmp_path, "U10707SX")
    write_upload_record(saved, "U10707SV", "saved_automatically", OrbeaWorkflowOptions.STAGES)
    calls = []
    def upload(match, root, **kwargs):
        calls.append((match.sku, kwargs["resume_completed"]))
        return OrbeaUploadResult(match, PimPreparationResult(match.sku, status=PimPreparationStatus.SAVED_AUTOMATICALLY))
    worker = OrbeaUploadWorker(lambda: SimpleNamespace(upload_and_save=upload),
        OrbeaUploadService.load_local_packages(tmp_path), None, retry_unfinished=True)
    worker.run()
    assert calls == [("U10707SX", True)]


def test_finished_status_prioritises_the_actual_batch_stop_over_the_first_failure(screen, tmp_path):
    make_package(tmp_path)
    panel = screen._upload_panel
    panel.load_packages(tmp_path)
    match = panel.matches[0]
    panel.results = [OrbeaUploadResult(match, PimPreparationResult(match.sku, status=PimPreparationStatus.FAILED, error="First row disappeared"))]
    panel.worker = SimpleNamespace(matches=(match,), _batch_error="Batch stopped because recovery did not finish", isRunning=lambda: False)
    panel._finished()
    assert "Reason: Batch stopped because recovery did not finish" in panel.status.text()


def test_worker_preflight_leaves_all_products_unprocessed_and_saved_progress_intact(tmp_path):
    folder = make_package(tmp_path, "U10707SV")
    make_package(tmp_path, "U10707SX")
    write_upload_record(folder, "U10707SV", "failed", ("description_source", "save"))
    saved_bytes = (folder / "orbea-upload-result.json").read_bytes()
    calls, summaries, errors = [], [], []
    def preflight(matches, **kwargs):
        assert len(matches) == 2
        raise RuntimeError("PIMBO has unsaved changes: Orbea AVANT H50")
    worker = OrbeaUploadWorker(lambda: SimpleNamespace(prepare_upload_batch=preflight,
        upload_and_save=lambda *args, **kwargs: calls.append(args)),
        OrbeaUploadService.load_local_packages(tmp_path), None, retry_unfinished=True)
    worker.completed.connect(summaries.append)
    worker.failed.connect(errors.append)
    worker.run()
    assert calls == [] and worker.results == []
    assert summaries[0]["failed"] == 0 and summaries[0]["unprocessed"] == 2
    assert "Orbea AVANT H50" in summaries[0]["error"] and errors == [summaries[0]["error"]]
    assert (folder / "orbea-upload-result.json").read_bytes() == saved_bytes


def test_worker_stop_includes_the_product_error_and_recovery_failure(tmp_path):
    make_package(tmp_path, "U10707SV")
    make_package(tmp_path, "U10707SX")
    matches = OrbeaUploadService.load_local_packages(tmp_path)
    calls, summaries = [], []
    def upload(match, root, **kwargs):
        calls.append(match.sku)
        return OrbeaUploadResult(match, PimPreparationResult(match.sku,
            status=PimPreparationStatus.FAILED, error="Product family could not be read"))
    def recover(*args, **kwargs):
        raise RuntimeError("The open PIMBO product changed while resetting")
    worker = OrbeaUploadWorker(lambda: SimpleNamespace(upload_and_save=upload,
        recover_after_failed_upload=recover), matches, None)
    worker.completed.connect(summaries.append)
    worker.run()
    assert calls == [matches[0].sku] and summaries[0]["unprocessed"] == 1
    assert "Product family could not be read" in summaries[0]["error"]
    assert "product changed while resetting" in summaries[0]["error"]


def test_activity_records_a_preflight_error_with_no_product_failures(screen, tmp_path):
    make_package(tmp_path)
    panel = screen._upload_panel
    panel.load_packages(tmp_path)
    records = []
    screen.main.operation_tracker = SimpleNamespace(finish=lambda *args, **kwargs: records.append((args, kwargs)))
    panel._operation_id = "operation-id"
    panel.results = []
    panel.worker = SimpleNamespace(matches=panel.matches, _stop_requested=False,
        _batch_error="Save or discard changes to Orbea AVANT H50", isRunning=lambda: False)
    panel._finish_activity()
    assert records[0][1]["summary"]["failed"] == 0
    assert records[0][1]["error_summary"] == panel.worker._batch_error


def test_upload_excel_export_after_reopening_folder_includes_pending_products(screen, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    folder = make_package(tmp_path)
    make_package(tmp_path, "U10707SX")
    (folder / "orbea-upload-result.json").write_text(json.dumps({"preparation": {"product_code": "U10707SV", "product_id": "p1",
        "status": "no_changes", "photo_upload": {"action": "skipped_existing", "existing_photos": 1}},
        "selected_stages": ["product_photos", "save"], "completed_stages": ["product_photos"]}))
    panel = screen._upload_panel
    panel._report_batch_error = "An older unrelated batch failed"
    panel.load_packages(tmp_path)
    assert panel._report_batch_error == ""
    assert panel.export_results.isEnabled() and screen.main.driver is None
    path = tmp_path / "upload.xlsx"
    monkeypatch.setattr("GUI_Qt.orbea.upload.QFileDialog.getSaveFileName", lambda *args: (str(path), "Excel (*.xlsx)"))
    panel._export_results()
    workbook = load_workbook(path)
    assert workbook["Results"].max_row == 3
    assert workbook["Needs checking"].max_row == 2
    assert workbook["Results"].cell(2, 7).value == "Skipped: real photos already present"
    assert workbook["Results"].cell(3, 4).value == "Unprocessed"
    workbook.close()
    panel._busy = True
    panel.update_state()
    assert not panel.export_results.isEnabled()
    panel._busy = False


def test_all_brand_history_export_includes_detailed_photo_results(screen, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from GUI_Qt.screens.FullHistoryScreen import FullHistoryScreen
    widget = QWidget()
    widget.filtered_history = [{"batch_id": None, "brand": "KROSS", "product_code": "SKU", "status": "saved_manually",
        "duration_seconds": 5, "features_uploaded": 0, "images_uploaded": 2, "error_message": "", "failed_stage": "",
        "url_or_code": "https://kross.pl/bike", "processed_at": "2026-10-07",
        "details_json": json.dumps({"pim_preparation": {"status": "saved_manually", "photo_upload": {
            "action": "uploaded", "uploaded_photos": 2, "placeholders_removed": 1}}})}]
    widget._record_type_label = lambda batch_id: "Single upload"
    path = tmp_path / "history.xlsx"
    monkeypatch.setattr("GUI_Qt.screens.FullHistoryScreen.QFileDialog.getSaveFileName", lambda *args: (str(path), "Excel (*.xlsx)"))
    success, errors = [], []
    monkeypatch.setattr("GUI_Qt.screens.FullHistoryScreen.InfoBar.success", lambda **kwargs: success.append(kwargs))
    monkeypatch.setattr("GUI_Qt.screens.FullHistoryScreen.InfoBar.error", lambda **kwargs: errors.append(kwargs))
    FullHistoryScreen._export_to_excel(widget)
    assert success and not errors
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Upload History", "Summary", "Results", "Needs checking"]
    assert workbook["Results"].cell(2, 8).value == 1
    assert workbook["Results"].cell(2, 11).value == "Yes"
    assert workbook["Needs checking"].max_row == 1
    workbook.close()
    widget.deleteLater()


@pytest.mark.parametrize("attribute,button_name", (
    ("_description_worker", "_description_start_btn"),
    ("_photo_worker", "_photo_start_btn"),
    ("_table_image_worker", "_table_image_start_btn"),
))
def test_direct_jobs_lock_upload_inputs_and_keep_stop_disabled_once_requested(screen, attribute, button_name):
    worker = SimpleNamespace(isRunning=lambda: True, _stop_requested=False)
    setattr(screen, attribute, worker)
    try:
        screen._update_action_states()
        assert not screen._upload_panel.isEnabled()
        assert getattr(screen, button_name).isEnabled()
        assert not screen._start_btn.isEnabled()
        assert not screen._load_collection_btn.isEnabled()
        assert all(not widget.isEnabled() for group in (
            screen._config_widgets, screen._description_config_widgets,
            screen._photo_config_widgets, screen._table_image_config_widgets,
        ) for widget in group)
        worker._stop_requested = True
        screen._update_action_states()
        assert not getattr(screen, button_name).isEnabled()
    finally:
        setattr(screen, attribute, None)
        screen._update_action_states()
    assert screen._upload_panel.isEnabled()


def test_completed_thread_keeps_job_ownership_until_queued_results_are_delivered(screen, tmp_path):
    from GUI_Qt.orbea.workers import OrbeaDescriptionWorker
    screen._description_urls_edit.setPlainText("https://www.orbea.com/en-be/m/ordu")
    screen._description_output_edit.setText(str(tmp_path))
    screen.workflow_controller.description_service_factory = lambda: SimpleNamespace(
        run=lambda *args, **kwargs: {"output_dir": tmp_path, "files": (), "failures": ()}
    )
    screen._on_description_start_stop()
    worker = screen._description_worker
    assert isinstance(worker, OrbeaDescriptionWorker)
    assert worker.wait(3000)
    assert not worker.isRunning()
    assert screen.is_running()
    screen._on_photo_start_stop()
    assert screen._photo_worker is None
    QApplication.instance().processEvents()
    assert screen._description_worker is None
    assert not screen.is_running()
    assert "complete" in screen._description_status_label.text().lower()


@pytest.mark.parametrize("cancelled", (False, True))
def test_description_failures_do_not_show_unqualified_success(screen, tmp_path, cancelled):
    warnings = []
    screen._warn = lambda *args: warnings.append(args)
    screen._on_description_result({
        "output_dir": tmp_path, "files": (tmp_path / "saved.txt",),
        "succeeded": 1, "failures": ("failed URL",), "cancelled": cancelled,
    })
    text = screen._description_status_label.text().lower()
    assert ("stopped" if cancelled else "failed") in text
    assert "1 saved" in screen._description_progress_label.text().lower()
    assert "1 failed" in screen._description_progress_label.text().lower()
    assert bool(warnings) is not cancelled


def test_direct_worker_start_failure_restores_the_controls(screen, tmp_path):
    def fail_tracking(*args, **kwargs):
        raise RuntimeError("worker registration failed")
    errors = []
    screen.main.track_worker = fail_tracking
    screen._error = lambda *args: errors.append(args)
    screen._description_urls_edit.setPlainText("https://www.orbea.com/en-be/m/ordu")
    screen._description_output_edit.setText(str(tmp_path))
    screen._on_description_start_stop()
    assert screen._description_worker is None and not screen.is_running()
    assert screen._description_start_btn.isEnabled()
    assert screen._upload_panel.isEnabled()
    assert errors and "worker registration failed" in errors[0][1]
