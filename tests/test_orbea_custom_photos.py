from __future__ import annotations

import base64
import json
from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from PIL import Image
from selenium.common.exceptions import ElementClickInterceptedException

from tests.test_orbea_colour_packages import PRODUCTS, URL, Website, Photos, photo_session, records, run
from tools.orbea_automation.custom_photos import (
    DESIGN_CONTROL_SCRIPT, GALLERY_METADATA_SCRIPT, GALLERY_SCRIPT, RELOAD_VIEWS_SCRIPT,
    CustomDesignAlreadySelected, CustomDesignNotReady, capture_default_design, save_default_design,
)
from tools.orbea_automation.models import CancellationToken, OrbeaRunConfig, RunCancelled
from tools.orbea_automation.upload import OrbeaUploadService
from tools.orbea_automation.website import OrbeaWebsiteClient


def snapshot(colour=(30, 150, 240, 255), views=('side', 'front')):
    image = Image.new('RGBA', (640, 360), colour)
    image.paste((0, 0, 0, 255), (100, 100, 550, 250))
    stream = BytesIO(); image.save(stream, format='PNG')
    data_url = 'data:image/png;base64,' + base64.b64encode(stream.getvalue()).decode()
    return {'asset_hash': 'default-design', 'template_id': '277', 'template_type': 'essential', 'images': [
        {'view': view, 'width': 640, 'height': 360, 'data_url': data_url} for view in views]}


class CustomWebsite(Website):
    def __init__(self, *, error=None):
        super().__init__()
        self.custom_calls = []
        self.custom_error = error

    def read_default_custom_photos(self, url, output_dir, **kwargs):
        self.custom_calls.append(url)
        if self.custom_error:
            raise self.custom_error
        return save_default_design(snapshot(), output_dir, **kwargs)


CUSTOM_PRODUCTS = (
    ('U21005CC', 'Orbea ALMA H20 Custom'),
    ('U21009CC', 'Orbea ALMA H20 2027 Custom'),
)


def test_custom_and_named_colours_share_one_model_visit_and_keep_separate_images(tmp_path):
    result, website, photos = run(tmp_path, (*CUSTOM_PRODUCTS, PRODUCTS[0]), website=CustomWebsite(), resume=False)
    assert result.completed
    assert website.fetches == [URL] and website.custom_calls == [URL]
    assert photos.calls == [('R6',)]
    rows = records(result)
    for row in rows[:2]:
        stage = row['collection_stages']['photos']
        assert stage['colour_code'] == 'Custom' and stage['colour'] == 'Your Design (default)'
        assert stage['pimbo_colour'] == 'Custom' and stage['source'] == 'your_design_default'
        folder = result.run_dir / row['local_folder']
        assert folder.parent.name == 'U210TTCC' and folder.name == row['sku']
        candidates = OrbeaUploadService.local_photo_candidates(folder)
        assert {p.name for p in candidates} == {'Custom_side.png', 'Custom_front.png'}
        with Image.open(candidates[0]) as image:
            assert image.getpixel((0, 0)) == (30, 150, 240, 255)
    assert rows[2]['collection_stages']['photos']['source'] == 'published_colour'
    assert rows[2]['collection_stages']['photos']['colour_code'] == 'R6'
    from openpyxl import load_workbook
    book = load_workbook(result.workbook_path, data_only=True)
    sheet = book['Collected Products']
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index('Photo colour') + 1).value == 'Your Design (default)'
    book.close()


def test_custom_only_does_not_require_or_download_factory_colour_templates(tmp_path):
    website = CustomWebsite()
    website.product = replace(website.product, page_html='<main><h1>ALMA H20</h1></main>')
    result, website, photos = run(tmp_path, CUSTOM_PRODUCTS, website=website, resume=False)
    assert result.completed and website.custom_calls == [URL]
    assert photos.calls == [] and photos._external_session.calls == []


def test_custom_failure_does_not_fail_named_colours_or_substitute_them(tmp_path):
    result, website, photos = run(tmp_path, (*CUSTOM_PRODUCTS, PRODUCTS[0]),
        website=CustomWebsite(error=ValueError('Your Design option missing')), resume=False)
    assert not result.completed and website.custom_calls == [URL] and photos.calls == [('R6',)]
    rows = records(result)
    assert [row['collection_stages']['photos']['status'] for row in rows] == ['error', 'error', 'downloaded']
    for row in rows[:2]:
        assert OrbeaUploadService.local_photo_candidates(result.run_dir / row['local_folder']) == ()
        assert 'Your Design option missing' in ' | '.join(row['collection_errors'])


def test_factory_template_failure_does_not_fail_custom_gallery(tmp_path):
    website = CustomWebsite()
    website.product = replace(website.product, page_html='<main>No factory template</main>')
    result, website, photos = run(tmp_path, (*CUSTOM_PRODUCTS, PRODUCTS[0]), website=website, resume=False)
    assert not result.completed and website.custom_calls == [URL] and photos.calls == []
    assert [row['collection_stages']['photos']['status'] for row in records(result)] == ['downloaded', 'downloaded', 'error']


def test_resume_skips_complete_customs_and_retries_only_failed_customs(tmp_path):
    result, _, _ = run(tmp_path, (*CUSTOM_PRODUCTS, PRODUCTS[0]),
        website=CustomWebsite(error=ValueError('Gallery did not load')), resume=False)
    rows = records(result)
    named_file = result.run_dir / rows[2]['collection_stages']['photos']['files'][0]
    original = named_file.read_bytes(), named_file.stat().st_mtime_ns
    settings = OrbeaRunConfig(None, tmp_path / 'runs', collect_product_data=True, download_images=False,
        download_product_photos=True, resume_run_dir=result.run_dir)
    retried, website, photos = run(tmp_path, settings=settings, website=CustomWebsite(), download_missing=True)
    assert retried.completed and website.custom_calls == [URL] and photos.calls == []
    assert original == (named_file.read_bytes(), named_file.stat().st_mtime_ns)
    completed, website, photos = run(tmp_path, settings=settings, website=CustomWebsite(), download_missing=True)
    assert completed.completed and website.fetches == [] and website.custom_calls == [] and photos.calls == []


def test_custom_cancellation_resumes_using_saved_named_photo_renders(tmp_path):
    result, _, _ = run(tmp_path, (PRODUCTS[0], *CUSTOM_PRODUCTS),
        website=CustomWebsite(error=RunCancelled('Stopped')), resume=False)
    assert result.cancelled
    settings = OrbeaRunConfig(None, tmp_path / 'runs', collect_product_data=True, download_images=False,
        download_product_photos=True, resume_run_dir=result.run_dir)
    resumed, website, photos = run(tmp_path, settings=settings, website=CustomWebsite(), download_missing=True)
    assert resumed.completed and website.custom_calls == [URL]
    # Cancellation happened before assigning the first row, but shared cached
    # factory renders can be reused without downloading their layers again.
    assert len(OrbeaUploadService.load_local_packages(resumed.run_dir / 'products')) == 3


class DesignBrowser:
    def __init__(self, *, frames=None, selected=False, available=True, intercepted=False):
        self.selected, self.available, self.intercepted = selected, available, intercepted
        self.frames = list(frames or [snapshot()])
        self.clicks, self.reads, self.exports, self.reloads = 0, 0, 0, 0
        self.input = object()
        self.control = SimpleNamespace(click=self.click)
        self.current_frame = None

    def click(self):
        if self.intercepted:
            self.intercepted = False
            raise ElementClickInterceptedException()
        self.clicks += 1
        self.selected = True

    def execute_script(self, script, *args):
        if script == DESIGN_CONTROL_SCRIPT:
            return {'input': self.input, 'control': self.control} if self.available else None
        if script == GALLERY_METADATA_SCRIPT:
            assert args == (self.input,) and self.selected
            self.reads += 1
            self.current_frame = self.frames.pop(0) if len(self.frames) > 1 else self.frames[0]
            if not self.current_frame or self.current_frame['template_type'] == 'standard':
                return None
            return {**self.current_frame, 'images': None, 'views': [
                {'view': image['view'], 'width': image['width'], 'height': image['height'],
                 'layers': [{'id': image['view'], 'url': image['data_url']}]} for image in self.current_frame['images']]}
        if script == GALLERY_SCRIPT:
            assert args == (self.input,) and self.selected
            self.exports += 1
            return self.current_frame
        if script == 'return arguments[0].checked === true;':
            return self.selected
        if script == 'arguments[0].click();':
            assert args == (self.control,)
            self.click(); return None
        raise AssertionError('Unexpected browser action: ' + script)

    def execute_async_script(self, script):
        assert script == RELOAD_VIEWS_SCRIPT
        self.reloads += 1
        self.frames = [snapshot()]
        return True


def test_your_design_is_clicked_once_and_still_loading_factory_views_are_rejected():
    factory = {**snapshot((170, 170, 170, 255)), 'template_type': 'standard'}
    default = snapshot()
    browser = DesignBrowser(frames=[factory, None, default])
    captured = capture_default_design(browser, timeout=3, settle_seconds=0)
    assert captured == default and browser.clicks == 1 and browser.reads == 3
    assert browser.exports == 1 and browser.reloads == 0


def test_default_views_must_be_stable_after_rendering_changes():
    default = snapshot()
    browser = DesignBrowser(frames=[snapshot((255, 30, 0, 255)), default, default])
    ticks = iter([0, .5, 2])
    with patch('tools.orbea_automation.custom_photos.time', SimpleNamespace(monotonic=lambda: next(ticks))):
        assert capture_default_design(browser, timeout=3, settle_seconds=1) == default
    assert browser.clicks == 1 and browser.reads == 3 and browser.exports == 1


def test_slow_or_missing_layers_are_rendered_again_without_customising_paint():
    browser = DesignBrowser(frames=[None])
    assert capture_default_design(browser, timeout=.01, settle_seconds=0) == snapshot()
    assert browser.clicks == browser.exports == browser.reloads == 1


def test_intercepted_label_click_uses_only_the_same_design_control():
    browser = DesignBrowser(intercepted=True)
    assert capture_default_design(browser, settle_seconds=0) == snapshot()
    assert browser.clicks == 1


def test_previously_customised_design_is_not_saved_as_a_default():
    browser = DesignBrowser(selected=True)
    with pytest.raises(ValueError, match='already selected'):
        capture_default_design(browser)
    assert browser.clicks == browser.reads == 0


def test_missing_design_option_is_an_error():
    browser = DesignBrowser(available=False)
    with pytest.raises(ValueError, match='no visible Your Design'):
        capture_default_design(browser, timeout=.01)
    assert browser.clicks == browser.reads == 0


def test_stop_while_waiting_for_custom_views_propagates():
    browser = DesignBrowser()
    calls = 0
    def check():
        nonlocal calls
        calls += 1
        if calls == 3: raise RunCancelled('Stopped')
    with pytest.raises(RunCancelled):
        capture_default_design(browser, check=check)
    assert browser.clicks == 1


def test_design_pngs_are_full_size_and_shared_in_the_image_store(tmp_path):
    result = save_default_design(snapshot(), tmp_path / 'model/Custom', asset_root=tmp_path / '.orbea-assets')
    assert result.variant_files == {'Custom': result.files} and len(result.files) == 2
    with Image.open(result.files[0]) as image: assert image.size == (640, 360)
    assert len(list((tmp_path / '.orbea-assets/images').glob('*.png'))) == 1
    manifest = json.loads((result.product_dir / 'download_manifest.json').read_text())
    assert manifest['source'] == 'your_design_default'


@pytest.mark.parametrize('invalid', ['blank', 'duplicate', 'not_png', 'missing', 'dimensions'])
def test_incomplete_or_invalid_custom_views_write_no_photos(tmp_path, invalid):
    data = snapshot()
    if invalid == 'duplicate': data['images'][1]['view'] = 'side'
    if invalid == 'not_png': data['images'][1]['data_url'] = 'https://other.example/image.png'
    if invalid == 'missing': data['images'] = []
    if invalid == 'dimensions': data['images'][1]['width'] = 10
    if invalid == 'blank':
        image = Image.new('RGBA', (640, 360), (255, 255, 255, 255))
        stream = BytesIO(); image.save(stream, format='PNG')
        data['images'][1]['data_url'] = 'data:image/png;base64,' + base64.b64encode(stream.getvalue()).decode()
    with pytest.raises(ValueError): save_default_design(data, tmp_path / 'Custom')
    assert list(tmp_path.rglob('*.png')) == []


def test_website_client_reuses_current_model_visit_for_custom_capture(tmp_path):
    browser = DesignBrowser()
    browser.current_url = URL
    browser.set_page_load_timeout = lambda seconds: None
    client = OrbeaWebsiteClient(lambda: browser, CancellationToken())
    with patch.object(client, '_page_soup'), patch.object(client, '_read_product_page') as navigate, \
         patch('tools.orbea_automation.custom_photos.capture_default_design', return_value=snapshot()) as capture:
        result = client.read_default_custom_photos(URL, tmp_path / 'Custom')
    navigate.assert_not_called()
    assert capture.call_args.args == (browser,) and len(result.files) == 2


def test_browser_scripts_use_the_real_checkbox_and_composite_all_loaded_layers(tmp_path):
    # This fixture models the public Orbea toggle and renderer captured from
    # ALMA M-LTD. Execute the production scripts to catch DOM lookup, readiness
    # and complete-bike layer ordering mistakes missed by Python adapters.
    import shutil
    import subprocess
    from pathlib import Path
    from bs4 import BeautifulSoup
    public = BeautifulSoup((Path(__file__).parent / 'fixtures/orbea_custom_toggle.html').read_text(encoding='utf-8'), 'html.parser')
    toggle = public.find('input', id='frame-options')
    assert toggle['type'] == 'checkbox'
    captions = [public.find('span', attrs={'x-ref': key}).get_text(' ', strip=True)
                for key in ('option-unchecked', 'option-checked')]
    assert captions == ['Our selection', 'You design']
    node = shutil.which('node')
    if not node: pytest.skip('Node is needed only for the browser-script fixture')
    fixture = r"""
const assert = require('node:assert/strict');
const scripts = JSON.parse(process.argv[2]);
let input = {type: 'checkbox', owned: true, checked: true, getClientRects: () => [{}]};
let directInput = input, labels = [], canvases = [];
let template = {id: 277, hash: 'custom-template', type: 'essential', views: [
    {type: 'front', order: 2, status: 'published'}, {type: 'side', order: 1, status: 'published'}]};
const standard = {id: 433, hash: 'standard-template', type: 'standard'};
const root = {contains: node => node.owned,
    querySelector: () => directInput,
    querySelectorAll: selector => selector === 'label' ? labels :
        canvases.filter(canvas => canvas.id.startsWith(selector.match(/id\^="(.+)"/)[1]))};
const state = {templates: [template, standard], template: 277, currentTemplate: template, customization: {}, initialize: true,
    isStandardTemplate: false, canvasesReady: true,
    imagesCanvas: {'277': {side: {60: 'frame.png', 85: 'wheels.png', base: 'base.png'},
        front: {60: 'frame-front.png', base: 'base-front.png'}}}};
const document = {querySelector: () => root, getElementById: id => canvases.find(canvas => canvas.id === id),
    createElement: () => {const draws = []; return {getContext: () => ({
        drawImage(canvas) { draws.push(canvas.id); }, clearRect() {} }), toDataURL: () => draws.join(',')};}};
const window = {Alpine: {$data: () => state}};
const getComputedStyle = node => ({visibility: node.invisible ? 'hidden' : 'visible'});
const execute = (script, ...args) => new Function('document', 'window', 'getComputedStyle', 'args',
    'return (function() {' + script + '}).apply(null, args);')(document, window, getComputedStyle, args);
const label = (text, choice=input, extra={}) => ({innerText: text, control: choice,
    getClientRects: () => extra.hidden ? [] : [{}], querySelector: () => extra.caption ? {textContent: extra.caption} : null, ...extra});
labels = [label(scripts.captions.join(' '), input, {caption: scripts.captions[1]})];
assert.equal(execute(scripts.control).input, input);
assert.equal(execute(scripts.control).control, input);
directInput = null; assert.equal(execute(scripts.control).input, input); // combined label fallback
labels.push(label('Your Design', {...input}));
assert.throws(() => execute(scripts.control), /More than one/);
labels = [label('Gloss'), label('Paint'), label('Your Design', {type: 'button', owned: true})];
assert.equal(execute(scripts.control), null);
labels = [label('Your Design', input, {hidden: true})]; assert.equal(execute(scripts.control), null);
labels = [label('Your Design', {...input, disabled: true})]; assert.equal(execute(scripts.control), null);
directInput = input;
state.initialize = false; assert.equal(execute(scripts.control), null); state.initialize = true;
state.customization = null; assert.equal(execute(scripts.control), null); state.customization = {};
state.canvasesReady = false; assert.equal(execute(scripts.control), null); state.canvasesReady = true;
state.loadingTemplate = 277; assert.equal(execute(scripts.control), null); state.loadingTemplate = null;
const canvas = (id, extra={}) => ({id, owned: true, dataset: {drawn: '1'}, width: 4100, height: 2310, ...extra});
canvases = [canvas('canvas-433_side_base'), canvas('canvas-277_side_60'), canvas('canvas-277_side_85'),
    canvas('canvas-277_side_alt60'), canvas('canvas-277_side_base'),
    canvas('canvas-277_side_name1', {dataset: {drawn: '1', textZone: 'name1'}}),
    canvas('canvas-277_front_60', {hidden: true}), canvas('canvas-277_front_base', {hidden: true}),
    canvas('canvas-277_side_unpublished', {dataset: {}})];
let result = execute(scripts.gallery, input);
assert.deepEqual(result.images.map(image => image.view), ['side', 'front']);
assert.equal(result.images[0].data_url, 'canvas-277_side_60,canvas-277_side_85,canvas-277_side_base,canvas-277_side_name1');
assert.equal(result.images[1].data_url, 'canvas-277_front_60,canvas-277_front_base');
assert.equal(result.images[0].width, 4100);
input.checked = false; assert.equal(execute(scripts.gallery, input), null); input.checked = true;
state.template = 433; assert.equal(execute(scripts.gallery, input), null); state.template = 277;
state.isStandardTemplate = true; assert.equal(execute(scripts.gallery, input), null); state.isStandardTemplate = false;
for (const flag of ['loadingTemplate', 'loadingInspiration', 'loadingVarnish', 'componentUpdating']) {
    state[flag] = true; assert.equal(execute(scripts.metadata, input), null); state[flag] = null;
}
state.canvasesReady = false; assert.equal(execute(scripts.metadata, input), null); state.canvasesReady = true;
canvases[2].dataset.drawn = ''; assert.equal(execute(scripts.metadata, input), null); canvases[2].dataset.drawn = '1';
canvases[2].width = 300; assert.equal(execute(scripts.metadata, input), null); canvases[2].width = 4100;
canvases[4].dataset.drawn = ''; assert.equal(execute(scripts.metadata, input), null); canvases[4].dataset.drawn = '1';
canvases[2].owned = false; assert.equal(execute(scripts.metadata, input), null); canvases[2].owned = true;
state.templates = [template]; labels = []; directInput = null;
assert.equal(execute(scripts.control).default_only, true);
assert.equal(execute(scripts.gallery, null).images.length, 2);
state.templates = [template, standard]; assert.equal(execute(scripts.gallery, null), null);
template.views.push({type: 'detail', order: 3, status: 'published'});
assert.deepEqual(execute(scripts.gallery, input).unavailable, ['detail']);
console.log('Real toggle, complete layer export and readiness checks passed');
"""
    path = tmp_path / 'custom-gallery-fixture.js'
    path.write_text(fixture, encoding='utf-8')
    result = subprocess.run([node, str(path), json.dumps({'control': DESIGN_CONTROL_SCRIPT,
        'metadata': GALLERY_METADATA_SCRIPT, 'gallery': GALLERY_SCRIPT, 'captions': captions})],
        capture_output=True, text=True, timeout=15)
    assert result.returncode == 0, result.stdout + result.stderr


@pytest.mark.parametrize('first_error', [CustomDesignAlreadySelected('already selected'), CustomDesignNotReady('loading stalled')])
def test_website_client_reloads_prior_customisation_or_stalled_loading_once(tmp_path, first_error):
    browser = DesignBrowser(selected=True)
    browser.current_url = URL
    browser.set_page_load_timeout = lambda seconds: None
    client = OrbeaWebsiteClient(lambda: browser, CancellationToken())
    with patch.object(client, '_page_soup'), patch.object(client, '_read_product_page') as navigate, \
         patch('tools.orbea_automation.custom_photos.capture_default_design',
               side_effect=[first_error, snapshot()]) as capture:
        result = client.read_default_custom_photos(URL, tmp_path / 'Custom')
    navigate.assert_called_once_with(URL)
    assert capture.call_args_list[1].kwargs['allow_selected_default'] is True
    assert len(result.files) == 2



def test_a_second_stalled_custom_page_is_reported_without_another_reload_or_photo_files(tmp_path):
    browser = DesignBrowser()
    browser.current_url = URL
    browser.set_page_load_timeout = lambda seconds: None
    client = OrbeaWebsiteClient(lambda: browser, CancellationToken())
    with patch.object(client, '_page_soup'), patch.object(client, '_read_product_page') as navigate, \
         patch('tools.orbea_automation.custom_photos.capture_default_design',
               side_effect=CustomDesignNotReady('loading stalled')) as capture:
        with pytest.raises(CustomDesignNotReady):
            client.read_default_custom_photos(URL, tmp_path / 'Custom')
    navigate.assert_called_once_with(URL)
    assert capture.call_count == 2 and not list(tmp_path.rglob('*.png'))


def test_fresh_page_can_use_its_already_selected_default_without_changing_paint():
    browser = DesignBrowser(selected=True)
    assert capture_default_design(browser, allow_selected_default=True, settle_seconds=0) == snapshot()
    assert browser.clicks == 0 and browser.exports == 1


def test_unpublished_custom_view_is_in_manifest_and_result(tmp_path):
    data = {**snapshot(), 'unavailable': ['detail']}
    result = save_default_design(data, tmp_path / 'Custom')
    manifest = json.loads((result.product_dir / 'download_manifest.json').read_text())
    assert result.unavailable == ('Custom Your Design view not published: detail',)
    assert manifest['template_id'] == '277' and manifest['unavailable'] == list(result.unavailable)


@pytest.mark.parametrize('title', [
    'Orbea OIZ M-PRO 2027 / Custom / kalnų (MTB) dviratis (29")',
    'Orbea OIZ M-PRO 2027 | CUSTOM | mountain bike',
    'Orbea OIZ M-PRO 2027 Custom (Matt) / mountain bike',
])
def test_custom_category_suffix_uses_design_gallery_and_keeps_spalva_manual(tmp_path, title):
    from tools.orbea_automation.specifications import parse_orbea_name, build_orbea_specification_plan
    model, colour = parse_orbea_name(title, source_model='OIZ M-PRO')
    assert model == 'Orbea OIZ M-PRO 2027' and colour == 'Custom'
    values = dict(build_orbea_specification_plan(title, ('M', 'S'), 'Frame: Carbon\nColor: Blue',
        product=SimpleNamespace(name='OIZ M-PRO')).values)
    assert 'Spalva' not in values and values['Galimi rėmo dydžiai'] == 'S, M'
    website = CustomWebsite()
    website.product = replace(website.product, name='OIZ M-PRO')
    result, website, photos = run(tmp_path, (('U21009CM', title),), website=website, resume=False)
    assert result.completed and website.custom_calls == [URL] and photos.calls == []
    assert records(result)[0]['collection_stages']['photos']['pimbo_colour'] == 'Custom'


def test_custom_word_inside_a_named_colour_is_not_the_custom_marker():
    from tools.orbea_automation.specifications import parse_orbea_name
    assert parse_orbea_name('Orbea OIZ M-PRO 2027 Custom Grey', source_model='OIZ M-PRO')[1] == 'Custom Grey'
    assert parse_orbea_name('Orbea OIZ M-PRO 2027 Frozen Concrete/Titanium', source_model='OIZ M-PRO')[1] == 'Frozen Concrete/Titanium'
