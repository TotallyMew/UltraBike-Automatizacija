"""Open an Orbea report as a saved collection without scanning or URL lookup."""
from __future__ import annotations

import json
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

from .catalogue import normalize_code
from .checkpoint import CHECKPOINT_NAME, CHECKPOINT_VERSION, RunCheckpoint, create_run_directory, saved_run_config
from .collection import stage_complete
from .models import OrbeaRunConfig, OrbeaRunResult
from .report import write_report
from .utils import canonicalize_url


def _sheet_rows(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, ())]
    for values in rows:
        yield dict(zip(headers, values))


def _read_report(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Matches" not in workbook.sheetnames:
            raise ValueError("Choose an Orbea report with a Matches sheet containing Variant SKU and Orbea URL")
        details = {}
        if "Summary" in workbook.sheetnames:
            details = {str(row[0]): row[1] for row in workbook["Summary"].iter_rows(values_only=True)
                       if len(row) > 1 and row[0]}
        raw = {}
        if "Raw Scan" in workbook.sheetnames:
            for row in _sheet_rows(workbook["Raw Scan"]):
                raw[(normalize_code(row.get("Variant SKU")), canonicalize_url(row.get("Orbea URL")))] = row
        matches = []
        seen = {}
        for index, row in enumerate(_sheet_rows(workbook["Matches"]), 2):
            sku = normalize_code(row.get("Variant SKU"))
            url = str(row.get("Orbea URL") or "").strip()
            if not sku and not url:
                continue
            parsed = urlsplit(url)
            host = (parsed.hostname or "").lower()
            if not sku or parsed.scheme not in {"http", "https"} or not (host == "orbea.com" or host.endswith(".orbea.com")) or parsed.username:
                raise ValueError(f"Matches row {index} needs a product code and a valid Orbea product URL")
            key = canonicalize_url(url)
            if sku in seen:
                if seen[sku] != key:
                    raise ValueError(f"The report has different Orbea links for {sku}; check that row before downloading")
                continue
            seen[sku] = key
            source = raw.get((sku, key), {})
            product_id = str(source.get("Product ID") or "")
            matches.append({"row_key": product_id or sku, "product_id": product_id,
                "sku": sku, "title": str(row.get("Pimbo Product") or source.get("Pimbo Product") or ""),
                "product_url": str(source.get("Pimbo URL") or ""), "status": "code_match",
                "catalogue_url": url, "catalogue_model": str(source.get("Catalogue Model") or ""),
                "catalogue_code": str(row.get("Catalogue Code") or ""),
                "match_method": "Saved Orbea Excel link", "website_lookup_status": "done"})
        if not matches:
            raise ValueError("This Excel has no matched Orbea product links")
        return matches, details
    finally:
        workbook.close()


def _pairs(rows):
    return {(normalize_code(row.get("sku")), canonicalize_url(row.get("catalogue_url")))
            for row in rows if row.get("status") == "code_match" and row.get("catalogue_url")}


def open_saved_collection(path, output_root, *, browser_name="chrome"):
    """Reuse a report's collection, or import its links into a new local run."""
    path = Path(path).resolve()
    matches, details = _read_report(path)
    candidates = [path.parent]
    if details.get("Collection folder"):
        candidates.append(Path(str(details["Collection folder"])).expanduser())
    run_name = str(details.get("Run") or "")
    if run_name and Path(run_name).name == run_name:
        candidates.append(Path(output_root) / run_name)
    for folder in dict.fromkeys(candidates):
        checkpoint_path = folder / CHECKPOINT_NAME
        if not checkpoint_path.is_file():
            continue
        try:
            data = json.loads(checkpoint_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if not isinstance(data, dict) or data.get("version") != CHECKPOINT_VERSION or _pairs(data.get("results", [])) != _pairs(matches):
            continue
        config = saved_run_config(folder, browser_name=browser_name)
        checkpoint = RunCheckpoint(checkpoint_path, data, resumed=True)
        break
    else:
        config = OrbeaRunConfig(None, Path(output_root), collect_product_data=True,
            download_images=False, download_product_photos=False, download_description=False,
            download_specifications=False, browser_name=browser_name, downloads_only=True)
        checkpoint = RunCheckpoint.create(create_run_directory(config.output_root), config)
        for row in matches:
            checkpoint.upsert_result(row)
        checkpoint.data.update(scan_completed=True, website_lookup_completed=True,
            images_completed=True, product_data_completed=True, product_photos_completed=True,
            saved_downloads_only=True, imported_report=str(path))
        checkpoint.mark_completed()
        write_report(checkpoint)
        config = saved_run_config(checkpoint.run_dir, browser_name=browser_name)
    workbook_path = checkpoint.workbook_path if checkpoint.workbook_path.is_file() else path
    result = OrbeaRunResult(checkpoint.run_dir, workbook_path, checkpoint.path,
        checkpoint.manifest_path, bool(checkpoint.data.get("completed")),
        bool(checkpoint.data.get("cancelled")), True, checkpoint.counts())
    return config, result


def collection_download_counts(run_dir):
    data = json.loads((Path(run_dir) / CHECKPOINT_NAME).read_text(encoding="utf-8"))
    rows = [row for row in data.get("results", []) if row.get("status") == "code_match" and row.get("catalogue_url")]
    return {"total": len(rows), **{stage: sum(stage_complete(stage,
        row.get("collection_stages", {}).get(stage, {}), run_dir) for row in rows)
        for stage in ("photos", "description", "specifications", "tables")}}
