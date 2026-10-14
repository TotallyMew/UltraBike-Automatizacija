"""Persist operator-supplied products in the ordinary Orbea collection."""
from __future__ import annotations

import json
import re
import shutil
import uuid
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .catalogue import normalize_code
from .checkpoint import CHECKPOINT_NAME, RunCheckpoint, atomic_write_json, create_run_directory, saved_run_config, utc_now
from .models import OrbeaRunConfig, OrbeaRunResult
from .photos import normalize_orbea_product_url
from .report import write_report
from .selected_downloads import resolve_product_rows
from .service import PIMBO_AUTOMATION_LOCK
from .website import WEBSITE_LOOKUP_VERSION, template_code


@dataclass(frozen=True)
class ManualProduct:
    sku: str
    name: str
    url: str
    model_code: str = ""

    def __post_init__(self):
        sku = str(self.sku or "").strip().upper()
        if not re.fullmatch(r"[A-Z][A-Z0-9]{3}(?:[A-Z0-9]{4})?", sku) or sku.endswith("TTCC") or not template_code(sku):
            raise ValueError("Enter the Pimbo product SKU, including size and colour when present; TTCC is a model code")
        code = str(self.model_code or template_code(sku)).strip().upper()
        if len(code) == 4:
            code += "TTCC"
        if code != template_code(sku):
            raise ValueError("The model code must match the first four characters of the product SKU")
        name = str(self.name or "").strip()
        if not name:
            raise ValueError("Enter the product name, including its colour for photo matching")
        object.__setattr__(self, "sku", sku)
        object.__setattr__(self, "model_code", code)
        object.__setattr__(self, "name", name)
        object.__setattr__(self, "url", normalize_orbea_product_url(self.url))


@dataclass(frozen=True)
class ManualProductSave:
    result: OrbeaRunResult
    source_root: Path
    row_key: str
    added: bool
    report_error: str = ""


@dataclass(frozen=True)
class ExcelProductRow:
    product: ManualProduct
    location: str


@dataclass(frozen=True)
class ManualProductsExcel:
    rows: tuple[ExcelProductRow, ...]
    issues: tuple[str, ...] = ()
    duplicates: int = 0


@dataclass(frozen=True)
class ManualProductsImport:
    saved: tuple[ManualProductSave, ...]
    issues: tuple[str, ...] = ()
    duplicates: int = 0

    @property
    def added(self):
        return sum(item.added for item in self.saved)

    @property
    def updated(self):
        return len(self.saved) - self.added


def read_manual_products_excel(path):
    """Read A/B/C/E product rows, or named columns, without changing the input."""
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook["Matches"] if "Matches" in workbook.sheetnames else workbook.active
        columns = {"sku": 0, "name": 1, "model_code": 2, "url": 4}
        aliases = {
            "sku": {"variant sku", "sku", "product code", "code", "prekės kodas", "prekes kodas"},
            "name": {"pimbo product", "product name", "product", "name", "prekės pavadinimas", "pavadinimas"},
            "model_code": {"catalogue code", "model code", "ttcc", "modelio kodas"},
            "url": {"orbea url", "orbea link", "url", "product link", "orbea nuoroda", "nuoroda"},
        }
        first = True
        rows, issues, duplicates, conflicts = {}, [], 0, set()
        for number, cells in enumerate(sheet.iter_rows(), 1):
            if not any(cell.value is not None or cell.hyperlink for cell in cells):
                continue
            if first:
                first = False
                headers = {re.sub(r"[\s_-]+", " ", str(cell.value or "").strip().casefold()): index
                           for index, cell in enumerate(cells)}
                named = {field: next((headers[key] for key in names if key in headers), None)
                         for field, names in aliases.items()}
                if named["sku"] is not None and named["name"] is not None and named["url"] is not None:
                    columns = named
                    continue

            def value(field):
                index = columns[field]
                if index is None or index >= len(cells):
                    return ""
                cell = cells[index]
                return cell.hyperlink.target if field == "url" and cell.hyperlink and cell.hyperlink.target else cell.value or ""

            values = {field: value(field) for field in columns}
            if not any(values.values()):  # Column D and unrelated columns are ignored.
                continue
            location = f"{sheet.title} row {number}"
            try:
                product = ManualProduct(**values)
            except ValueError as error:
                issues.append(f"{location}: {error}")
                continue
            if product.sku in rows:
                previous = rows[product.sku]
                if previous.product == product:
                    duplicates += 1
                else:
                    conflicts.add(product.sku)
                    issues.append(f"{location}: conflicting entries for {product.sku}; check this row and {previous.location}")
                continue
            rows[product.sku] = ExcelProductRow(product, location)
        return ManualProductsExcel(tuple(row for sku, row in rows.items() if sku not in conflicts), tuple(issues), duplicates)
    finally:
        workbook.close()


def _invalidate_package(checkpoint, row, product):
    """Retain old files/history, but prevent their use with a corrected source."""
    stages = {stage: {"status": "pending", "files": [], "note": ""}
              for stage in ("photos", "description", "specifications", "tables")}
    if row.get("local_folder"):
        folder = (checkpoint.run_dir / row["local_folder"]).resolve()
        if not folder.is_relative_to(checkpoint.run_dir.resolve()):
            raise ValueError("The saved product folder is outside the collection")
        metadata = folder / "orbea-product.json"
        if metadata.is_file():
            saved = json.loads(metadata.read_text(encoding="utf-8-sig"))
            if normalize_code(saved.get("sku")) != product.sku or str(saved.get("pimbo_product_id") or "") != str(row.get("product_id") or ""):
                raise ValueError("The product package belongs to a different saved row")
            saved.update(orbea_url=product.url, pimbo_product_name=product.name,
                         orbea_product_name=product.name, status="not_collected",
                         stages=stages, errors=[], warnings=[])
            atomic_write_json(metadata, saved)
        history = folder / "orbea-upload-result.json"
        if history.is_file():
            saved = json.loads(history.read_text(encoding="utf-8-sig"))
            saved.update(completed_stages=[], source_changed_at=utc_now())
            if isinstance(saved.get("preparation"), dict):
                saved["source_previous_status"] = saved["preparation"].get("status")
                saved["preparation"]["status"] = "ready_for_review"
            atomic_write_json(history, saved)
    row.update(collection_stages=stages, collection_status="not_collected", collection_errors=[])


class _ManualProductWriter:
    """Reuse loaded checkpoints and export once per run for an Excel batch."""
    def __init__(self, output_root, source_root, browser_name):
        self.root = Path(source_root).resolve() if source_root else None
        if self.root is not None and not self.root.is_dir():
            raise ValueError("The displayed collection folder is missing")
        if self.root is None and not str(output_root).strip():
            raise ValueError("Choose the Orbea output folder first")
        self.output_root = self.root or Path(output_root)
        self.browser_name = browser_name
        self.checkpoints, self.owners, self.records = {}, {}, []
        self.new_checkpoint = None
        if self.root is not None:
            paths = [self.root / CHECKPOINT_NAME] if (self.root / CHECKPOINT_NAME).is_file() else sorted(self.root.rglob(CHECKPOINT_NAME))
            for path in paths:
                if not path.resolve().is_relative_to(self.root):
                    raise ValueError("A saved collection is outside the displayed folder")
                saved_run_config(path.parent)
                checkpoint = RunCheckpoint(path, json.loads(path.read_text(encoding="utf-8-sig")), resumed=True)
                self.checkpoints[path] = checkpoint
                for row in checkpoint.results:
                    self.owners.setdefault(normalize_code(row.get("sku")), []).append((checkpoint, row))

    def apply(self, product, *, target=None):
        if not isinstance(product, ManualProduct):
            raise TypeError("A validated ManualProduct is required")
        checkpoint = None
        row = None
        if target is not None:
            config, owned = resolve_product_rows((target,))[0]
            if owned["sku"] != product.sku:
                raise ValueError("An existing product's SKU cannot be changed; add a new product instead")
            if self.root is None or not config.resume_run_dir.resolve().is_relative_to(self.root):
                raise ValueError("The product does not belong to the displayed collection")
            checkpoint = self.checkpoints[config.resume_run_dir / CHECKPOINT_NAME]
            row = next(item for item in checkpoint.results if item["row_key"] == owned["row_key"])
        else:
            owners = self.owners.get(product.sku, [])
            if len(owners) > 1:
                raise ValueError("This SKU occurs in multiple saved rows; select a row and use Edit product")
            if owners:
                checkpoint, row = owners[0]
            elif self.root is not None:
                checkpoint = self.checkpoints.get(self.root / CHECKPOINT_NAME)
        added = row is None
        if checkpoint is None:
            if self.new_checkpoint is None:
                config = OrbeaRunConfig(None, self.output_root, browser_name=self.browser_name,
                    collect_product_data=True, downloads_only=True, download_images=False)
                self.new_checkpoint = RunCheckpoint.create(create_run_directory(config.output_root), config)
                self.new_checkpoint.data.update(scan_completed=True, saved_downloads_only=True, website_lookup_completed=True)
            checkpoint = self.new_checkpoint
        if row is None:
            row = {"row_key": "manual:" + uuid.uuid4().hex, "product_id": "", "product_url": "", "sku": product.sku}
        source_changed = row.get("catalogue_url") != product.url or row.get("title") != product.name
        if source_changed:
            _invalidate_package(checkpoint, row, product)
        if added:
            checkpoint.results.append(row)
            self.owners.setdefault(product.sku, []).append((checkpoint, row))
        row.update(title=product.name, catalogue_code=product.model_code,
                   catalogue_model=product.name if source_changed else row.get("catalogue_model") or product.name,
                   catalogue_url=product.url, status="code_match", match_method="Manual product link",
                   manual_product=True, manual_updated_at=utc_now(), website_lookup_status="done",
                   website_lookup_version=WEBSITE_LOOKUP_VERSION, website_lookup_source="manual", note="")
        row.pop("website_validation_error", None)
        if checkpoint.data.get("photo_scan_skus") is not None and product.sku not in checkpoint.data["photo_scan_skus"]:
            checkpoint.data["photo_scan_skus"].append(product.sku)
        checkpoint.data.update(completed=False, product_data_completed=False, product_photos_completed=False,
                               images_completed=False)
        self.records.append((checkpoint, row["row_key"], added))

    def finish(self, *, source_path=None):
        results = {}
        for checkpoint, _key, _added in self.records:
            if checkpoint.path in results:
                continue
            checkpoint.save()
            report_error = ""
            try:
                if source_path is not None and checkpoint.workbook_path.resolve() == Path(source_path).resolve():
                    # Keep even a collection's own report intact when used as input.
                    destination = checkpoint.run_dir / "orbea_matches_imported.xlsx"
                    number = 2
                    while destination.exists():
                        destination = checkpoint.run_dir / f"orbea_matches_imported_{number}.xlsx"
                        number += 1
                    shutil.copyfile(source_path, destination)
                    checkpoint.data["report_filename"] = destination.name
                    checkpoint.save()
                write_report(checkpoint)
            except Exception as error:
                report_error = str(error)
            result = OrbeaRunResult(checkpoint.run_dir, checkpoint.workbook_path, checkpoint.path,
                checkpoint.manifest_path, False, False, checkpoint.resumed, checkpoint.counts())
            results[checkpoint.path] = result, report_error
        return tuple(ManualProductSave(results[checkpoint.path][0], self.root or checkpoint.run_dir,
                     key, added, results[checkpoint.path][1]) for checkpoint, key, added in self.records)


def save_manual_product(product, output_root, *, source_root=None, target=None, browser_name="chrome"):
    """Add or update one full SKU, preserving existing row/Pimbo identities."""
    if not PIMBO_AUTOMATION_LOCK.acquire(blocking=False):
        raise RuntimeError("Another Orbea operation is running")
    try:
        writer = _ManualProductWriter(output_root, source_root, browser_name)
        writer.apply(product, target=target)
        return writer.finish()[0]
    finally:
        PIMBO_AUTOMATION_LOCK.release()


def import_manual_products_excel(path, output_root, *, source_root=None, browser_name="chrome"):
    """Append valid Excel rows; report invalid/ambiguous rows without replacing the list."""
    excel = read_manual_products_excel(path)
    if not excel.rows:
        raise ValueError("No products could be imported. Expected A: SKU, B: Name, C: TTCC code, E: Orbea URL.\n" + "\n".join(excel.issues))
    if not PIMBO_AUTOMATION_LOCK.acquire(blocking=False):
        raise RuntimeError("Another Orbea operation is running")
    try:
        writer = _ManualProductWriter(output_root, source_root, browser_name)
        issues = list(excel.issues)
        for entry in excel.rows:
            try:
                writer.apply(entry.product)
            except ValueError as error:
                issues.append(f"{entry.location}: {error}")
            except OSError as error:
                issues.append(f"{entry.location}: import stopped: {error}. Later rows were not imported.")
                break
        return ManualProductsImport(writer.finish(source_path=path), tuple(issues), excel.duplicates)
    finally:
        PIMBO_AUTOMATION_LOCK.release()
