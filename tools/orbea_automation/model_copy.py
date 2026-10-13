"""Copy saved, manually prepared Pimbo models to their exact TTCC siblings."""
from __future__ import annotations

import hashlib
import html
import json
import re
from dataclasses import asdict, dataclass, field, replace
from pathlib import Path
from typing import Callable

from Managers.PimboProductEditor import PimAutomationError, PimPreparationStatus, PimboProductEditor
from Managers.PimboChoices import choices_match
from tools.kross_automation.service import KrossPimboClient
from tools.supplier_upload import PIMBO_PRODUCTS_URL, SupplierUploadWorkflow, normalize_sku
from .checkpoint import atomic_write_json, utc_now
from .website import template_code
from .colour_attributes import load_colour_mapping

COPY_SPECIFICATIONS = ("Modelis", "Ratų dydis", "Rėmo medžiaga", "Rėmas")
CHECKPOINT_NAME = "model-copy-checkpoint.json"


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def source_title_parts(title: str) -> tuple[str, str, str]:
    parts = re.split(r"\s+/\s+", clean(title))
    if len(parts) != 3 or not all(parts):
        raise ValueError('Source name must be: model year / colour / bicycle type (wheel size)')
    model, colour, suffix = parts
    if not re.match(r"^Orbea\s+.+\s+(?:19|20)\d{2}$", model, re.I):
        raise ValueError("Source name must include Orbea, its model and a final four-digit year")
    return model, colour, suffix


def target_colour(title: str, model: str) -> str:
    title = clean(title)
    base = re.sub(r"\s+(?:19|20)\d{2}$", "", model)
    segments = re.split(r"\s+/\s+", title)
    if len(segments) > 1:
        if len(segments) not in (2, 3):
            raise ValueError("Target name has ambiguous colour separators")
        target_model = re.sub(r"\s+(?:19|20)\d{2}$", "", segments[0])
        if target_model.casefold() != base.casefold():
            raise ValueError("Target name does not match the source model")
        colour = segments[1]
    else:
        match = re.match(re.escape(base) + r"\s+(?:(?:19|20)\d{2}\s+)?(.+)$", title, re.I)
        if not match:
            raise ValueError("Cannot reliably separate the target model and colour")
        colour = match.group(1)
    if not clean(colour) or re.fullmatch(r"(?:19|20)\d{2}", clean(colour)):
        raise ValueError("Target colour is empty")
    return clean(colour)


def lacquer_finish(colour: str) -> str:
    values = ("Gloss" if word.casefold().startswith("gloss") else "Matt"
              for word in re.findall(r"\b(?:glossy?|matte?)\b", colour, re.I))
    return " - ".join(dict.fromkeys(values))


@dataclass
class OrbeaCopyRow:
    sku: str
    title: str
    source: bool = False


@dataclass
class OrbeaModelTemplate:
    ttcc: str
    source_sku: str
    product_id: str
    product_url: str
    version: int
    model: str
    suffix: str
    title: str
    categories: list[dict]
    attributes: dict[str, list[str]]
    specifications: dict[str, str]
    description_html: str

    @property
    def fingerprint(self):
        return hashlib.sha256(json.dumps(asdict(self), sort_keys=True, ensure_ascii=False).encode()).hexdigest()


@dataclass
class OrbeaCopyGroup:
    ttcc: str
    rows: list[OrbeaCopyRow]
    source_sku: str = ""
    status: str = "pending"
    error: str = ""
    template: OrbeaModelTemplate | None = None
    source_product_id: str = ""
    source_selected_in_app: bool = False
    workbook_error: str = ""

    @property
    def targets(self):
        return [row for row in self.rows if row.sku != self.source_sku]


@dataclass
class OrbeaModelCopyResult:
    ttcc: str
    source_sku: str
    sku: str
    status: str
    product_id: str = ""
    product_url: str = ""
    title: str = ""
    failed_stage: str = ""
    error: str = ""
    template_fingerprint: str = ""
    checked_at: str = field(default_factory=utc_now)
    source_product_id: str = ""
    source_version: int | None = None
    colour_name: str = ""
    colour_attribute: list[str] = field(default_factory=list)
    colour_mapping_fingerprint: str = ""


@dataclass
class OrbeaCopyRunResult:
    groups: list[OrbeaCopyGroup]
    results: list[OrbeaModelCopyResult]
    checkpoint_path: Path
    cancelled: bool = False
    error: str = ""


def load_copy_workbook(path: Path | str) -> list[OrbeaCopyGroup]:
    from openpyxl import load_workbook
    path = Path(path)
    if not path.is_file() or path.suffix.casefold() != ".xlsx":
        raise ValueError("Choose an existing .xlsx file")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["Matches"] if "Matches" in book.sheetnames else book.active
        rows = iter(sheet.iter_rows(values_only=True))
        headers = [clean(value).casefold() for value in next(rows, ())]
        required = ("variant sku", "pimbo product", "catalogue code")
        missing = [name for name in required if headers.count(name) != 1]
        if headers.count("šaltinis") > 1:
            missing.append("šaltinis")
        if missing:
            raise ValueError("Missing or duplicate Excel columns: " + ", ".join(missing))
        indexes = {name: headers.index(name) for name in required}
        if "šaltinis" in headers:
            indexes["šaltinis"] = headers.index("šaltinis")
        groups = {}
        for row in rows:
            if not any(value is not None for value in row):
                continue
            values = {name: clean(row[index]) if index < len(row) else "" for name, index in indexes.items()}
            sku = normalize_sku(values["variant sku"])
            ttcc = values["catalogue code"].upper()
            group = groups.setdefault(ttcc, OrbeaCopyGroup(ttcc, []))
            marker = values.get("šaltinis", "").casefold()
            if marker not in ("", "taip"):
                group.error = "Šaltinis must be TAIP or empty"
            if (not re.fullmatch(r"[A-Z][A-Z0-9]{7}", sku)
                    or not re.fullmatch(r"[A-Z][A-Z0-9]{3}TTCC", ttcc) or template_code(sku) != ttcc):
                group.error = f"SKU/TTCC mismatch: {sku} / {ttcc}"
            if any(item.sku == sku for item in group.rows):
                group.error = f"Duplicate SKU in Excel: {sku}"
            group.rows.append(OrbeaCopyRow(sku, values["pimbo product"], marker == "taip"))
        for group in groups.values():
            sources = [row for row in group.rows if row.source]
            if len(sources) > 1:
                group.error = "Multiple sources marked TAIP in one TTCC group"
            group.source_sku = sources[0].sku if len(sources) == 1 else ""
            group.status = ("error" if group.error else "unmarked" if not sources else
                            "no_targets" if not group.targets else "pending")
            group.workbook_error = group.error
        return list(groups.values())
    finally:
        book.close()


class OrbeaModelCopyService(SupplierUploadWorkflow):
    """A separate checkpointed workflow; Stop takes effect between products."""

    def __init__(self, driver, *, editor_factory=PimboProductEditor, pimbo_client_factory=KrossPimboClient):
        self.pimbo_driver = driver
        self.editor_factory = editor_factory
        self.pimbo_client_factory = pimbo_client_factory
        self._stopped = False
        self.colour_mapping = load_colour_mapping()

    def set_colour_mapping(self, mapping):
        self.colour_mapping = mapping

    def _colour_mapping_changed(self, result, template):
        if result.status != "saved" or not template or not result.title:
            return False
        try:
            fingerprint = self.colour_mapping.fingerprint_for(target_colour(result.title, template.model))
        except ValueError:
            return False
        return bool(fingerprint and fingerprint != result.colour_mapping_fingerprint)

    def cancel(self):
        self._stopped = True

    def _stop(self, cancellation):
        return self._stopped or bool(cancellation and cancellation.is_cancelled())

    @staticmethod
    def checkpoint_path(workbook):
        path = Path(workbook).resolve()
        return path.parent / (path.stem + "_model_copy") / CHECKPOINT_NAME

    def _resolve(self, sku):
        match = self.pimbo_client_factory(self.pimbo_driver).find_by_variant_sku(sku)
        if match.status != "pimbo_found" or not match.pimbo_product_id:
            raise PimAutomationError(match.note or f"Pimbo variant {sku} was not found")
        editor = self.editor_factory(self.pimbo_driver)
        editor.wait_ready()
        if editor.product_id != match.pimbo_product_id:
            raise PimAutomationError("Resolved Pimbo product identity changed")
        return match, editor

    def _leave_read_only_product(self, editor, product_id):
        # Sources and duplicate targets are only inspected. Discard editor
        # defaults only while the verified product is still open.
        if editor.product_id != product_id:
            raise PimAutomationError("Inspected product identity changed; browser recovery is unsafe")
        self.pimbo_driver.get(PIMBO_PRODUCTS_URL)

    def _unverified_product_open(self):
        return bool(self._pimbo_product_id(getattr(self.pimbo_driver, "current_url", "")))

    @staticmethod
    def _scope(groups):
        return [(group.ttcc, group.source_sku, [row.sku for row in group.rows]) for group in groups]

    @classmethod
    def _write(cls, path, workbook, groups, results):
        atomic_write_json(path, {"version": 1, "workbook": str(Path(workbook).resolve()),
                                "groups": [asdict(group) for group in groups],
                                "results": [asdict(result) for result in results], "updated_at": utc_now()})

    @classmethod
    def load_checkpoint(cls, workbook, *, allow_scope_change=False):
        path = cls.checkpoint_path(workbook)
        data = json.loads(path.read_text(encoding="utf-8"))
        if (not isinstance(data, dict) or data.get("version") != 1
                or data.get("workbook") != str(Path(workbook).resolve()) or not isinstance(data.get("groups"), list)):
            raise ValueError("Model-copy checkpoint belongs to another workbook or version")
        groups = []
        for value in data["groups"]:
            value = dict(value)
            value["rows"] = [OrbeaCopyRow(**row) for row in value["rows"]]
            value["template"] = OrbeaModelTemplate(**value["template"]) if value.get("template") else None
            groups.append(OrbeaCopyGroup(**value))
        for group in groups:
            template = group.template
            if template and (template.ttcc != group.ttcc or template.source_sku != group.source_sku
                             or not template.product_id or not isinstance(template.version, int) or template.version < 0):
                raise ValueError("Invalid saved model template")
            if template and (set(template.specifications) - set(COPY_SPECIFICATIONS)
                             or any(name.casefold() == "spalva" for name in template.attributes)):
                raise ValueError("Saved template contains fields outside model-copy scope")
        results = [OrbeaModelCopyResult(**item) for item in data.get("results", [])]
        allowed = {(group.ttcc, group.source_sku, row.sku) for group in groups for row in group.targets}
        if any((item.ttcc, item.source_sku, item.sku) not in allowed for item in results):
            raise ValueError("Checkpoint results are outside the Excel scope")
        fingerprints = {group.ttcc: group.template.fingerprint for group in groups if group.template}
        if any(item.template_fingerprint != fingerprints.get(item.ttcc) for item in results):
            raise ValueError("Checkpoint result belongs to another source snapshot")
        current = load_copy_workbook(workbook)
        previous = {group.ttcc: group for group in groups}
        for group in current:
            old = previous.get(group.ttcc)
            if (old and old.source_selected_in_app and not group.workbook_error
                    and (not old.source_sku or any(row.sku == old.source_sku for row in group.rows))):
                cls._apply_source_choice(group, old.source_sku)
        invalid_templates = {group.ttcc for group in groups if group.template}
        changed = (cls._scope(groups) != cls._scope(current)
                   or any(group.workbook_error and group.ttcc in invalid_templates for group in current)
                   or any(group.ttcc in previous and group.workbook_error != previous[group.ttcc].workbook_error for group in current))
        if changed:
            if not allow_scope_change:
                raise ValueError("Excel sources or target rows changed. Read sources again.")
            # Keep only snapshots whose full source/target scope is unchanged.
            # New or changed groups must be read again before they can be copied.
            groups = [previous[group.ttcc] if group.ttcc in previous and not group.workbook_error
                      and cls._scope([group]) == cls._scope([previous[group.ttcc]]) else group for group in current]
            valid = {group.ttcc: group.template.fingerprint for group in groups if group.template}
            results = [item for item in results if valid.get(item.ttcc) == item.template_fingerprint]
        return OrbeaCopyRunResult(groups, results, path)

    @staticmethod
    def _apply_source_choice(group, sku):
        if group.workbook_error:
            raise ValueError(group.workbook_error)
        if sku and not any(row.sku == sku for row in group.rows):
            raise ValueError(f"Source {sku} is not in TTCC group {group.ttcc}")
        group.source_sku = sku
        for row in group.rows:
            row.source = row.sku == sku
        group.source_selected_in_app = True
        group.template, group.source_product_id, group.error = None, "", ""
        group.status = "unmarked" if not sku else "no_targets" if not group.targets else "pending"

    @classmethod
    def select_source(cls, workbook, ttcc, sku):
        path = cls.checkpoint_path(workbook)
        state = (cls.load_checkpoint(workbook, allow_scope_change=True) if path.is_file()
                 else OrbeaCopyRunResult(load_copy_workbook(workbook), [], path))
        group = next((group for group in state.groups if group.ttcc == ttcc), None)
        if group is None:
            raise ValueError(f"TTCC group {ttcc} is not in this workbook")
        if group.source_sku != sku or not group.source_selected_in_app:
            cls._apply_source_choice(group, sku)
            state.results = [item for item in state.results if item.ttcc != ttcc]
        cls._write(path, workbook, state.groups, state.results)
        return state

    def capture_templates(self, workbook, *, selected_ttcc=None, progress=None, log=None, cancellation=None, item_finished=None):
        groups = load_copy_workbook(workbook)
        path = self.checkpoint_path(workbook)
        self.prepare_upload_batch((None, None), progress=log)
        previous = []
        previous_groups = {}
        try:
            previous_state = self.load_checkpoint(workbook)
            previous = previous_state.results
            previous_groups = {group.ttcc: group for group in previous_state.groups}
        except (OSError, ValueError, KeyError, TypeError):
            pass
        for group in groups:
            old = previous_groups.get(group.ttcc)
            if old and old.source_selected_in_app and not group.workbook_error:
                self._apply_source_choice(group, old.source_sku)
        selected = set(selected_ttcc) if selected_ttcc is not None else {group.ttcc for group in groups}
        candidates = {group.ttcc for group in groups if group.status == "pending" and group.ttcc in selected}
        for group in groups:
            old = previous_groups.get(group.ttcc)
            if old and self._scope([group]) == self._scope([old]):
                group.template, group.source_product_id = old.template, old.source_product_id
                group.status, group.error = old.status, old.error
        # Re-reading can stop between sources. Keep confirmed results for groups
        # not inspected this time, rather than silently forgetting their saves.
        results = list(previous)
        batch_error = ""
        for group in groups:
            if self._stop(cancellation):
                break
            if group.ttcc not in candidates:
                continue
            results = [item for item in results if item.ttcc != group.ttcc]
            match, editor = None, None
            try:
                if progress:
                    progress(f"Reading source {group.source_sku} ({group.ttcc})")
                match, editor = self._resolve(group.source_sku)
                group.source_product_id = match.pimbo_product_id
                editor.switch_locale("lt")
                version = editor.current_version()
                title = editor.product_name()
                model, _colour, suffix = source_title_parts(title)
                description = editor.description_html(require_content=True)
                categories = editor.selected_categories()
                attributes = editor.read_attribute_values()
                specifications = {name: value for name in COPY_SPECIFICATIONS
                                  if (value := editor.specification_value(name))}
                if (editor.product_id != match.pimbo_product_id or version is None
                        or editor.current_version() != version):
                    raise PimAutomationError("Source version is unavailable or changed while reading")
                if not re.sub(r"<[^>]+>", "", html.unescape(description)).strip():
                    raise PimAutomationError("Source LT description is empty")
                if not categories:
                    raise PimAutomationError("Source has no selected categories")
                group.template = OrbeaModelTemplate(group.ttcc, group.source_sku, editor.product_id,
                    match.pimbo_product_url, version, model, suffix, title, categories,
                    {name: list(values) for name, values in attributes.items()
                     if name.casefold() != "spalva" and values}, specifications, description)
                group.status, group.error = "ready", ""
                # Source inspection is read-only. Return to the list after the
                # verified read instead of carrying editor defaults into a batch.
                self._leave_read_only_product(editor, match.pimbo_product_id)
                results.extend(item for item in previous if item.ttcc == group.ttcc
                               and item.template_fingerprint == group.template.fingerprint)
            except Exception as error:
                group.status, group.error = "error", str(error)
                group.template = None
                if log:
                    log(f"{group.ttcc}: {error}")
                try:
                    if editor is not None:
                        self._leave_read_only_product(editor, match.pimbo_product_id)
                    elif self._unverified_product_open():
                        raise PimAutomationError("Cannot safely leave an unverified source product")
                except Exception as recovery:
                    batch_error = f"{error}; recovery: {recovery}"
            self._write(path, workbook, groups, results)
            if item_finished:
                item_finished(group)
            if batch_error:
                break
        self._write(path, workbook, groups, results)
        return OrbeaCopyRunResult(groups, results, path, self._stop(cancellation), batch_error)

    def _verify_fields(self, editor, template, title, finish, colours):
        editor.switch_locale("lt")
        if editor.product_name() != title or not choices_match(editor.selected_categories(), template.categories, categories=True):
            raise PimAutomationError("Copied name or categories did not persist")
        attributes = editor.read_attribute_values()
        for name, values in template.attributes.items():
            if sorted(attributes.get(name, [])) != sorted(values):
                raise PimAutomationError(f"Attribute {name!r} did not persist")
        if sorted(attributes.get("Spalva", [])) != sorted(colours[0]):
            raise PimAutomationError("The target Spalva attribute did not retain its expected values")
        if editor.specification_value("Spalva") != colours[1]:
            raise PimAutomationError("The target Spalva specification changed; it must always be preserved")
        for name, value in {**template.specifications, "Lako užbaigimas": finish}.items():
            if editor.specification_value(name) != value:
                raise PimAutomationError(f"Specification {name!r} did not persist")
        if editor.normalize_description_html(editor.description_html()) != editor.normalize_description_html(template.description_html):
            raise PimAutomationError("Copied description HTML did not persist")

    def run(self, workbook, *, selected_ttcc=None, retry_failed=False, progress=None, log=None,
            cancellation=None, item_finished: Callable | None = None):
        state = self.load_checkpoint(workbook)
        groups, results, path = state.groups, state.results, state.checkpoint_path
        self.prepare_upload_batch((None, None), progress=log)
        selected = set(selected_ttcc) if selected_ttcc is not None else {group.ttcc for group in groups if group.template and group.status == "ready"}
        source_ids = {group.source_product_id for group in groups if group.source_product_id}
        source_ids.update(group.template.product_id for group in groups if group.template)
        templates = {group.ttcc: group.template for group in groups}
        seen = {item.product_id for item in results if item.status == "saved" and item.product_id
                and not self._colour_mapping_changed(item, templates.get(item.ttcc))}
        batch_error = ""
        for group in groups:
            template = group.template
            if group.ttcc not in selected or not template or group.status != "ready" or self._stop(cancellation):
                continue
            source, source_match = None, None
            try:
                source_match, source = self._resolve(template.source_sku)
                if source.product_id != template.product_id or source.current_version() != template.version:
                    raise PimAutomationError("Source changed. Read sources again before copying this group.")
                group.status, group.error = "ready", ""
                self._leave_read_only_product(source, source_match.pimbo_product_id)
            except Exception as error:
                group.status, group.error = "source_changed", str(error)
                self._write(path, workbook, groups, results)
                if log:
                    log(f"{group.ttcc}: {error}")
                try:
                    if source is not None:
                        self._leave_read_only_product(source, source_match.pimbo_product_id)
                    elif self._unverified_product_open():
                        raise PimAutomationError("Cannot safely leave an unverified source product")
                except Exception as recovery:
                    batch_error = f"{error}; recovery: {recovery}"
                    break
                continue
            for row in group.targets:
                if self._stop(cancellation):
                    break
                old = next((item for item in results if item.ttcc == group.ttcc and item.sku == row.sku), None)
                if old and old.status in {"saved", "duplicate", "source_product"} and not self._colour_mapping_changed(old, template):
                    continue
                if retry_failed and (old is None or old.status != "failed"):
                    continue
                stage, match = "resolve", None
                result = OrbeaModelCopyResult(group.ttcc, group.source_sku, row.sku, "failed",
                    template_fingerprint=template.fingerprint, source_product_id=template.product_id,
                    source_version=template.version)
                try:
                    if progress:
                        progress(f"Copying {row.sku} from {group.source_sku}")
                    match, editor = self._resolve(row.sku)
                    result.product_id, result.product_url = match.pimbo_product_id, match.pimbo_product_url
                    if result.product_id in source_ids:
                        result.status = "source_product"
                        self._leave_read_only_product(editor, result.product_id)
                    elif result.product_id in seen:
                        result.status = "duplicate"
                        self._leave_read_only_product(editor, result.product_id)
                    else:
                        seen.add(result.product_id)
                        stage = "draft"
                        base = replace(editor.begin(""), product_code=row.sku)
                        if base.status == PimPreparationStatus.BLOCKED_NON_DRAFT:
                            raise PimAutomationError(base.error or "Target is not Draft")
                        editor.switch_locale("lt")
                        stage = "title"
                        colour = target_colour(editor.product_name(), template.model)
                        title = f"{template.model} / {colour} / {template.suffix}"
                        finish = lacquer_finish(colour)
                        result.title = title
                        mapped_colour = self.colour_mapping.values_for(colour)
                        result.colour_name = colour
                        result.colour_attribute = list(mapped_colour)
                        result.colour_mapping_fingerprint = self.colour_mapping.fingerprint_for(colour)
                        colours = (list(mapped_colour) if mapped_colour else editor.read_attribute_values().get("Spalva", []),
                                   editor.specification_value("Spalva"))
                        if not mapped_colour and log:
                            log(f"{row.sku}: no Spalva mapping for {colour}; keeping the existing colour attribute")
                        stage = "categories"
                        editor.set_categories(template.categories)
                        editor.set_product_name(title)
                        stage = "attributes"
                        for name, values in template.attributes.items():
                            editor.set_attribute_values(name, values)
                        if mapped_colour:
                            stage = "colour_attribute"
                            editor.set_attribute_values("Spalva", list(mapped_colour))
                        stage = "specifications"
                        for name, value in {**template.specifications, "Lako užbaigimas": finish}.items():
                            if editor.set_specification(name, value) is None:
                                raise PimAutomationError(f"Specification {name!r} was not found")
                        stage = "description"
                        editor.set_description_html(template.description_html)
                        self._verify_fields(editor, template, title, finish, colours)
                        stage = "translation"
                        translated = editor.translate_lt_to_all(overwrite=True, allow_unchanged=True)
                        if not translated.success:
                            raise PimAutomationError(translated.detail or "Translation failed")
                        editor.switch_locale("lt")
                        self._verify_fields(editor, template, title, finish, colours)
                        translations = {locale: editor.localized_copy_snapshot(locale) for locale in ("en", "lv", "ee")}
                        editor.switch_locale("lt")
                        stage = "save"
                        if editor.is_dirty():
                            prepared = editor.finish(base, changed_fields=("model_copy",), ai_steps=(translated,))
                            saved = editor.save_and_verify(prepared)
                            if saved.status != PimPreparationStatus.SAVED_AUTOMATICALLY:
                                raise PimAutomationError(saved.error or "Save verification failed")
                        # Read persisted data after a reload, even when values were already equal.
                        stage = "verify"
                        self.pimbo_driver.get(result.product_url)
                        editor = self.editor_factory(self.pimbo_driver)
                        editor.wait_ready()
                        if editor.product_id != result.product_id:
                            raise PimAutomationError("Product changed during save verification")
                        self._verify_fields(editor, template, title, finish, colours)
                        for locale, expected in translations.items():
                            persisted = editor.localized_copy_snapshot(locale)
                            if (persisted.get("name") != expected.get("name") or
                                    editor.normalize_description_html(persisted.get("description")) !=
                                    editor.normalize_description_html(expected.get("description"))):
                                raise PimAutomationError(f"{locale.upper()} translated copy did not persist")
                        editor.switch_locale("lt")
                        result.status = "saved"
                        seen.add(result.product_id)
                        # This editor was freshly reloaded solely for read-back.
                        # Leave it now, before Pimbo defaults can be mistaken for
                        # manual unsaved edits by the next SKU navigation guard.
                        if editor.product_id != result.product_id:
                            raise PimAutomationError("Verified product changed before returning to Products")
                        self.pimbo_driver.get(PIMBO_PRODUCTS_URL)
                except Exception as error:
                    result.status = "failed"
                    result.failed_stage, result.error = stage, str(error)
                    if match:
                        try:
                            if not self.recover_after_failed_upload(match, progress=log):
                                raise PimAutomationError("Could not safely recover the verified failed product")
                        except Exception as recovery:
                            batch_error = f"{result.error}; recovery: {recovery}"
                    else:
                        # Resolution can leave an unverified product open. Never discard it.
                        if self._unverified_product_open():
                            batch_error = str(error)
                results = [item for item in results if not (item.ttcc == group.ttcc and item.sku == row.sku)] + [result]
                self._write(path, workbook, groups, results)
                if item_finished:
                    item_finished(result)
                if log:
                    log(f"{row.sku}: {result.status}" + (f" — {result.error}" if result.error else ""))
                if batch_error:
                    break
            if batch_error:
                break
        self._write(path, workbook, groups, results)
        return OrbeaCopyRunResult(groups, results, path, self._stop(cancellation), batch_error)


def export_copy_results(state: OrbeaCopyRunResult, destination):
    """Export separately from the operator's source workbook."""
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Copy results"
    sheet.append(("TTCC", "Source SKU", "Target SKU", "Status", "Product ID", "Name", "Failed stage", "Error", "Checked at UTC",
                  "Source product ID", "Source version", "Orbea colour", "Spalva attribute"))
    lookup = {(result.ttcc, result.sku): result for result in state.results}
    for group in state.groups:
        for row in group.targets:
            result = lookup.get((group.ttcc, row.sku))
            sheet.append((group.ttcc, group.source_sku, row.sku, result.status if result else group.status,
                          result.product_id if result else "", result.title if result else row.title,
                          result.failed_stage if result else "", result.error if result else group.error,
                          result.checked_at if result else "", result.source_product_id if result else group.source_product_id,
                          result.source_version if result else group.template.version if group.template else "",
                          result.colour_name if result else "", "; ".join(result.colour_attribute) if result else ""))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip("ABCDEFGHIJKLM", (18, 18, 18, 20, 30, 80, 20, 80, 25, 30, 18, 60, 25)):
        sheet.column_dimensions[column].width = width
    from openpyxl.styles import Font, PatternFill
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="263850")
    from Utilities.ExportFiles import atomic_save_workbook
    return atomic_save_workbook(book, destination)
