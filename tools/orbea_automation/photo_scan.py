"""Excel-driven, read-only checks of uploaded PIMBO product photos."""
from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, replace
from pathlib import Path

from openpyxl import load_workbook

from Managers.PimboProductEditor import PimboProductEditor
from tools.kross_automation.service import KrossPimboClient
from .catalogue import normalize_code
from .checkpoint import CHECKPOINT_NAME, RunCheckpoint, atomic_write_json, create_run_directory, saved_run_config, utc_now
from .models import CancellationToken, OrbeaRunConfig, OrbeaRunResult
from .photos import normalize_orbea_product_url
from .report import write_report
from .service import PIMBO_AUTOMATION_LOCK
from .upload import OrbeaUploadMatch


PHOTO_STATUSES = {"not_checked": "Not checked", "no_photos": "No photos",
                  "has_photos": "Has photos", "not_found": "Not found", "scan_error": "Scan error"}
CODE_HEADERS = ("variant sku", "sku", "product code", "code", "prekės kodas", "prekes kodas")


def _header(value):
    return re.sub(r"[\s_-]+", " ", str(value or "").strip().casefold())


def read_photo_scan_products(path):
    """Accept a saved Orbea report or an Excel sheet with a code/SKU column."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(dict.fromkeys([name for name in ("Pimbo Photos", "Matches") if name in workbook.sheetnames]
                                  + workbook.sheetnames))
        for name in names:
            values = workbook[name].iter_rows(values_only=True)
            for header_row, headers in enumerate(values, 1):
                if header_row > 20:
                    break
                lookup = {_header(value): index for index, value in enumerate(headers) if value}
                code_header = next((key for key in CODE_HEADERS if key in lookup), None)
                if code_header is None:
                    continue

                def cell(row, *keys):
                    return next((row[lookup[key]] for key in keys
                                 if key in lookup and lookup[key] < len(row) and row[lookup[key]] is not None), "")

                products = {}
                for row_number, row in enumerate(values, header_row + 1):
                    raw_code = cell(row, code_header)
                    if isinstance(raw_code, float) and raw_code.is_integer():
                        raw_code = int(raw_code)
                    sku = normalize_code(raw_code)
                    if not sku:
                        continue
                    url = str(cell(row, "orbea url", "orbea link", "product link") or "").strip()
                    if url:
                        try:
                            url = normalize_orbea_product_url(url)
                        except ValueError as error:
                            raise ValueError(f"{name} row {row_number}: {error}") from error
                    product = {"sku": sku, "title": str(cell(row, "pimbo product", "product", "product name", "name") or ""),
                               "catalogue_url": url, "row_key": sku, "product_id": "",
                               "status": "code_match" if url else "unmatched",
                               "match_method": "Excel product link" if url else "Excel code; Orbea lookup pending"}
                    if sku in products and products[sku]["catalogue_url"] and url and products[sku]["catalogue_url"] != url:
                        raise ValueError(f"The Excel has different Orbea links for {sku}")
                    if sku not in products or (url and not products[sku]["catalogue_url"]):
                        products[sku] = product
                if not products:
                    raise ValueError(f"The {name} sheet has no product codes to scan")
                return tuple(products.values()), name
        raise ValueError("Choose an Excel with a Variant SKU, SKU, Product Code or Code column")
    finally:
        workbook.close()


def load_photo_scan_collection(path, output_root, *, browser_name="chrome"):
    path = Path(path).resolve()
    products, sheet = read_photo_scan_products(path)
    checkpoint_path = path.parent / CHECKPOINT_NAME
    checkpoint = None
    if checkpoint_path.is_file():
        data = json.loads(checkpoint_path.read_text(encoding="utf-8"))
        available = {row.get("sku") for row in data.get("results", [])}
        if {row["sku"] for row in products} <= available:
            saved_run_config(path.parent, browser_name=browser_name)  # validate version
            for product in products:
                saved = [row for row in data.get("results", []) if row.get("sku") == product["sku"]]
                if len(saved) != 1:
                    raise ValueError(f"The saved collection has multiple products for {product['sku']}")
                if product["catalogue_url"]:
                    saved_url = saved[0].get("catalogue_url")
                    if saved_url and normalize_orbea_product_url(saved_url) != product["catalogue_url"]:
                        raise ValueError(f"The Excel Orbea link for {product['sku']} differs from the saved collection")
                    saved[0].update(catalogue_url=product["catalogue_url"], status="code_match")
            checkpoint = RunCheckpoint(checkpoint_path, data, resumed=True)
    if checkpoint is None and sheet == "Matches" and all(row["catalogue_url"] for row in products):
        from .saved_collection import open_saved_collection
        _config, result = open_saved_collection(path, output_root, browser_name=browser_name)
        checkpoint = RunCheckpoint(result.checkpoint_path,
            json.loads(result.checkpoint_path.read_text(encoding="utf-8")), resumed=True)
    if checkpoint is None:
        config = OrbeaRunConfig(None, Path(output_root), browser_name=browser_name,
            collect_product_data=True, download_images=False, download_product_photos=False,
            download_description=False, download_specifications=False, downloads_only=True)
        checkpoint = RunCheckpoint.create(create_run_directory(config.output_root), config)
        checkpoint.data["results"] = list(products)
        checkpoint.data.update(scan_completed=True, saved_downloads_only=True,
                               images_completed=True, product_data_completed=True, product_photos_completed=True,
                               website_lookup_completed=all(row["catalogue_url"] for row in products))
    checkpoint.data.update(photo_scan_source=str(path), photo_scan_skus=[row["sku"] for row in products])
    checkpoint.save()
    write_report(checkpoint)
    return saved_run_config(checkpoint.run_dir, browser_name=browser_name), OrbeaRunResult(
        checkpoint.run_dir, checkpoint.workbook_path, checkpoint.path, checkpoint.manifest_path,
        bool(checkpoint.data.get("completed")), False, checkpoint.resumed, checkpoint.counts())


def photo_scan_rows(root):
    path = Path(root) / CHECKPOINT_NAME
    if not path.is_file():
        return ()
    data = json.loads(path.read_text(encoding="utf-8"))
    skus = data.get("photo_scan_skus")
    return tuple(row for row in data.get("results", []) if row.get("sku") and (skus is None or row["sku"] in skus))


def photo_scan_matches(root, packages):
    """Join checkpoint identities to packages, including not-yet-collected rows."""
    indexed = {str(Path(match.local_folder).resolve()): match for match in packages}
    matches = []
    owned_folders = set()
    for run, row in saved_product_rows(root):
        folder = (run / row.get("local_folder", f"products/{row['sku']}")).resolve()
        if not folder.is_relative_to(run):
            raise ValueError("A saved product folder is outside the collection")
        owned_folders.add(str(folder))
        package = indexed.get(str(folder))
        if package and (package.sku != row["sku"] or package.pimbo_product_id != str(row.get("product_id") or "")):
            raise ValueError(f"{row['sku']}: the product package belongs to a different saved row")
        if package is not None:
            matches.append(replace(package, checkpoint_path=str(run / CHECKPOINT_NAME), row_key=row["row_key"]))
            continue
        matches.append(OrbeaUploadMatch(row["sku"], "not_collected", str(row.get("product_id") or ""),
            row.get("title", ""), row.get("product_url", ""), row.get("catalogue_model", ""),
            row.get("catalogue_url", ""), str(folder), row.get("note", ""),
            checkpoint_path=str(run / CHECKPOINT_NAME), row_key=row["row_key"]))
    # A report imported beside its checkpoint may deliberately contain only a
    # subset. Do not re-add excluded rows just because their packages exist.
    scopes = {}
    for folder, match in indexed.items():
        if folder in owned_folders:
            continue
        for parent in Path(folder).parents:
            checkpoint = parent / CHECKPOINT_NAME
            if checkpoint.is_file():
                if checkpoint not in scopes:
                    scopes[checkpoint] = json.loads(checkpoint.read_text(encoding="utf-8")).get("photo_scan_skus")
                scope = scopes[checkpoint]
                if scope is None or match.sku in scope:
                    matches.append(match)
                break
        else:
            matches.append(match)
    return tuple(matches)


def saved_product_rows(root):
    """Read a run or a folder of runs without combining products by SKU alone."""
    root = Path(root).resolve()
    paths = [root / CHECKPOINT_NAME] if (root / CHECKPOINT_NAME).is_file() else sorted(root.rglob(CHECKPOINT_NAME))
    rows = []
    for path in paths:
        run = path.parent.resolve()
        if not run.is_relative_to(root):
            raise ValueError("A saved collection is outside the selected folder")
        saved_run_config(run)  # validate the checkpoint format
        rows.extend((run, row) for row in photo_scan_rows(run) if row.get("row_key"))
    return tuple(rows)


@dataclass(frozen=True)
class PimboPhotoScanResult:
    sku: str
    status: str
    product_id: str = ""
    product_name: str = ""
    product_url: str = ""
    photos: int | None = None
    placeholders: int | None = None
    error: str = ""
    checked_at: str = ""
    checkpoint_path: str = ""


class PimboPhotoScanService:
    def __init__(self, driver, *, client_factory=KrossPimboClient, editor_factory=PimboProductEditor):
        self.driver = driver
        self.client_factory = client_factory
        self.editor_factory = editor_factory
        self._cancellation = CancellationToken()

    def cancel(self):
        self._cancellation.cancel()

    def run(self, run_dir, skus, *, cancellation=None, progress=None, log=None, item_finished=None):
        token = cancellation or self._cancellation
        self._cancellation = token
        path = Path(run_dir) / CHECKPOINT_NAME
        checkpoint = RunCheckpoint(path, json.loads(path.read_text(encoding="utf-8")), resumed=True)
        rows = {row["sku"]: row for row in checkpoint.results if row.get("sku")}
        skus = tuple(dict.fromkeys(skus))
        if not skus or not set(skus) <= rows.keys():
            raise ValueError("Select products from the loaded Excel or saved collection")
        for sku in skus:
            if sum(row.get("sku") == sku for row in checkpoint.results) != 1:
                raise ValueError(f"The saved collection has multiple products for {sku}")
        if not PIMBO_AUTOMATION_LOCK.acquire(blocking=False):
            raise RuntimeError("Another Orbea Pimbo operation is running")
        results = []
        try:
            client = self.client_factory(self.driver)
            for index, sku in enumerate(skus, 1):
                if token.is_cancelled():
                    break
                row = rows[sku]
                if progress:
                    progress(f"Checking uploaded Pimbo photos: {index}/{len(skus)} — {sku}")
                try:
                    found = client.find_by_variant_sku(sku)
                    if found.status != "pimbo_found":
                        result = PimboPhotoScanResult(sku, "not_found", error=found.note, checked_at=utc_now())
                    else:
                        if row.get("product_id") and row["product_id"] != found.pimbo_product_id:
                            raise ValueError("The saved product code now belongs to a different Pimbo product")
                        editor = self.editor_factory(self.driver)
                        counts = editor.inspect_product_photos(expected_product_id=found.pimbo_product_id)
                        result = PimboPhotoScanResult(sku, "has_photos" if counts["photos"] else "no_photos",
                            found.pimbo_product_id, found.pimbo_product_name, found.pimbo_product_url,
                            counts["photos"], counts["placeholders"], checked_at=utc_now())
                        row.update(product_id=result.product_id, title=result.product_name, product_url=result.product_url)
                        if row.get("local_folder"):
                            folder = (checkpoint.run_dir / row["local_folder"]).resolve()
                            if not folder.is_relative_to(checkpoint.run_dir.resolve()):
                                raise ValueError("The saved product folder is outside the collection")
                            metadata = folder / "orbea-product.json"
                            if metadata.is_file():
                                saved = json.loads(metadata.read_text(encoding="utf-8"))
                                if saved.get("sku") == sku and saved.get("pimbo_product_id") in ("", None, result.product_id):
                                    saved.update(pimbo_product_id=result.product_id, pimbo_product_name=result.product_name,
                                                 pimbo_url=result.product_url)
                                    atomic_write_json(metadata, saved)
                except Exception as error:
                    result = PimboPhotoScanResult(sku, "scan_error", error=str(error), checked_at=utc_now())
                result = replace(result, checkpoint_path=str(path.resolve()))
                row["pimbo_photo_scan"] = asdict(result)
                checkpoint.save()
                results.append(result)
                if log:
                    log(f"{sku}: {PHOTO_STATUSES[result.status]}" +
                        (f" ({result.photos} uploaded photos)" if result.photos is not None else "") +
                        (f" — {result.error}" if result.error else ""))
                if item_finished:
                    item_finished(result)
            checkpoint.data["photo_scan_cancelled"] = token.is_cancelled()
            checkpoint.save()
            return tuple(results)
        finally:
            try:
                write_report(checkpoint)
                if log and checkpoint.workbook_path.name != "orbea_matches.xlsx":
                    log(f"Excel report saved separately because the original was open or read-only: {checkpoint.workbook_path}")
            finally:
                PIMBO_AUTOMATION_LOCK.release()


def missing_photo_download_keys(run_dir, skus):
    rows = photo_scan_rows(run_dir)
    keys = tuple(row["row_key"] for row in rows if row.get("sku") in skus
                 and row.get("pimbo_photo_scan", {}).get("status") == "no_photos" and row.get("product_id"))
    if not keys:
        raise ValueError("Scan Pimbo first and select products marked No photos")
    return keys


def record_uploaded_photos(result):
    """Update a scan only after the upload has confirmed product attachment."""
    upload = result.preparation.photo_upload or {}
    count = upload.get("uploaded_photos", 0)
    if upload.get("action") != "uploaded" or not upload.get("persisted") or count <= 0:
        return None
    match = result.match
    folder = Path(match.local_folder).resolve()
    run_dir = next((parent for parent in folder.parents if (parent / CHECKPOINT_NAME).is_file()), None)
    if run_dir is None:
        return None
    path = run_dir / CHECKPOINT_NAME
    data = json.loads(path.read_text(encoding="utf-8"))
    for row in data.get("results", []):
        if (row.get("sku") == match.sku and "pimbo_photo_scan" in row
                and str(row.get("product_id") or "") == result.preparation.product_id
                and (run_dir / row.get("local_folder", f"products/{match.sku}")).resolve() == folder):
            state = asdict(PimboPhotoScanResult(match.sku, "has_photos", result.preparation.product_id,
                row.get("title", ""), row.get("product_url", ""), count, 0, checked_at=utc_now()))
            row["pimbo_photo_scan"] = state
            atomic_write_json(path, data)
            return state
    return None
