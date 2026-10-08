"""Excel reports for supplier uploads and the app's saved upload history."""
from __future__ import annotations

import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

STATUS_LABELS = {
    "saved_automatically": "Saved automatically", "saved_manually": "Saved manually",
    "no_changes": "No changes needed", "success": "Successful", "ready_for_review": "Needs Save / review",
    "blocked_non_draft": "Blocked: not Draft", "failed": "Failed", "error": "Failed",
    "discarded": "Discarded", "unprocessed": "Unprocessed", "not_ready": "Not ready",
}
PHOTO_LABELS = {
    "uploaded": "Uploaded", "skipped_existing": "Skipped: real photos already present",
    "missing_photos": "No replacement photos available", "failed": "Failed",
    "ready": "Checked; upload not completed", "checking": "Inspection not completed",
}
HEADERS = (
    "Product code", "Product", "Brand", "Status", "Needs checking", "Reason to check",
    "Photo result", "Placeholders removed", "Existing real photos", "Photos added to form",
    "Photo changes saved", "Selected steps", "Completed steps", "Changed fields", "Warnings",
    "Error", "Failed step", "Pimbo URL", "Pimbo product ID", "Source URL", "Local folder",
    "Finished at", "Upload type", "Duration (s)", "Features uploaded", "Images uploaded",
)
SUCCESS_STATUSES = {"saved_automatically", "saved_manually", "no_changes", "success"}


def result_record(result):
    return {"finished_at": datetime.now(timezone.utc).isoformat(),
            "preparation": result.preparation.to_dict(),
            "selected_stages": list(result.options.selected_stages),
            "completed_stages": list(result.completed_stages)}


def upload_result_row(match, record=None, *, brand=""):
    record = record or {}
    preparation = record.get("preparation") or {}
    photo = preparation.get("photo_upload") or {}
    status = preparation.get("status") or ("unprocessed" if getattr(match, "ready", True) else "not_ready")
    selected = record.get("selected_stages") or []
    completed = record.get("completed_stages") or []
    missing = set(selected) - set(completed)
    if status == "no_changes":
        missing.discard("save")
    warnings = list(preparation.get("warnings") or [])
    error = preparation.get("error") or ""
    photo_failed = photo.get("action") in {"failed", "missing_photos", "checking", "ready"}
    reasons = [error, *warnings]
    if photo_failed:
        reasons.append(photo.get("error") or PHOTO_LABELS.get(photo.get("action"), "Photo step did not complete"))
    if missing:
        reasons.append("Steps not completed: " + ", ".join(stage for stage in selected if stage in missing))
    if status not in SUCCESS_STATUSES:
        reasons.append(STATUS_LABELS.get(status, status))
    if not preparation and getattr(match, "note", ""):
        reasons.append(match.note)
    needs_check = status not in SUCCESS_STATUSES or bool(error or warnings or missing or photo_failed)
    photo_changed = bool(photo.get("uploaded_photos") or photo.get("placeholders_removed"))
    saved = status in {"saved_automatically", "saved_manually"} or "save" in completed
    photo_saved = ("Yes" if saved else "No — review required") if photo_changed else ("Not needed" if photo.get("action") == "skipped_existing" else "")
    return {
        "Product code": preparation.get("product_code") or getattr(match, "sku", ""),
        "Product": getattr(match, "pimbo_product_name", "") or getattr(match, "orbea_product_name", "") or getattr(match, "kross_product_name", ""),
        "Brand": brand, "Status": STATUS_LABELS.get(status, status), "Needs checking": "Yes" if needs_check else "No",
        "Reason to check": "\n".join(dict.fromkeys(item for item in reasons if item)),
        "Photo result": PHOTO_LABELS.get(photo.get("action"), "Not recorded" if "product_photos" in selected else "Not selected"),
        "Placeholders removed": photo.get("placeholders_removed", ""), "Existing real photos": photo.get("existing_photos", ""),
        "Photos added to form": photo.get("uploaded_photos", ""), "Photo changes saved": photo_saved,
        "Selected steps": ", ".join(selected), "Completed steps": ", ".join(completed),
        "Changed fields": ", ".join(preparation.get("changed_fields") or []), "Warnings": "\n".join(warnings),
        "Error": error, "Failed step": preparation.get("failed_stage") or "",
        "Pimbo URL": preparation.get("final_url") or getattr(match, "pimbo_product_url", ""),
        "Pimbo product ID": preparation.get("product_id") or getattr(match, "pimbo_product_id", ""),
        "Source URL": getattr(match, "orbea_url", "") or getattr(match, "kross_url", ""),
        "Local folder": getattr(match, "local_folder", ""), "Finished at": record.get("finished_at") or "",
    }


def history_upload_row(record):
    from types import SimpleNamespace
    record = dict(record)
    details = record.get("details_json") or {}
    if isinstance(details, str):
        try:
            details = json.loads(details)
        except (ValueError, TypeError):
            details = {}
    details = details if isinstance(details, dict) else {}
    preparation = dict(details.get("pim_preparation") or {})
    preparation.setdefault("status", record.get("status") or "unprocessed")
    preparation.setdefault("product_code", record.get("product_code") or "")
    preparation["error"] = preparation.get("error") or record.get("error_message") or ""
    preparation["failed_stage"] = preparation.get("failed_stage") or record.get("failed_stage") or ""
    match = SimpleNamespace(sku=record.get("product_code") or "", ready=True, note="")
    row = upload_result_row(match, {"preparation": preparation, "finished_at": record.get("processed_at")}, brand=record.get("brand") or "")
    row.update({"Upload type": record.get("batch_id") or "Single upload", "Duration (s)": record.get("duration_seconds"),
                "Features uploaded": record.get("features_uploaded") or 0, "Images uploaded": record.get("images_uploaded") or 0})
    row["Source URL"] = record.get("url_or_code") or ""
    if not preparation.get("photo_upload"):
        row["Photo result"] = "Not recorded"
    return row


def add_upload_results(workbook, rows, *, context=None):
    """Add complete results and a filtered review list without interpreting cell text as formulas."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    rows = list(rows)
    summary = workbook.create_sheet("Summary")
    counts = Counter(row.get("Status", "Unprocessed") for row in rows)
    for key, value in {"Exported at": datetime.now().isoformat(timespec="seconds"), "Products": len(rows),
                       "Needs checking": sum(row.get("Needs checking") == "Yes" for row in rows), **(context or {}), **counts}.items():
        summary.append([str(key), value])
    for title, selected_rows in (("Results", rows), ("Needs checking", [row for row in rows if row.get("Needs checking") == "Yes"])):
        sheet = workbook.create_sheet(title)
        sheet.append(HEADERS)
        for row in selected_rows:
            sheet.append([row.get(header, "") for header in HEADERS])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="253449")
            cell.font = Font(bold=True, color="FFFFFF")
        for index, header in enumerate(HEADERS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 42 if header in {"Reason to check", "Warnings", "Error", "Pimbo URL", "Source URL", "Local folder"} else 24
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 80
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_upload_results(path, rows, *, context=None):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_upload_results(workbook, rows, context=context)
    workbook.save(Path(path))


def read_supplier_result(match, brand):
    if not getattr(match, "local_folder", ""):
        return None
    try:
        record = json.loads((Path(match.local_folder) / f"{brand}-upload-result.json").read_text(encoding="utf-8-sig"))
        preparation = record["preparation"]
        if preparation.get("product_code", "").upper() != match.sku.upper():
            return None
        if getattr(match, "pimbo_product_id", "") and preparation.get("product_id") != match.pimbo_product_id:
            return None
        return record
    except (OSError, ValueError, KeyError, TypeError, AttributeError):
        return None
